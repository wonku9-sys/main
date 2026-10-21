"""Bounded, data-only Excel catalogue readers. Never opens Excel or runs macros."""
from pathlib import Path
import struct
import math
from decimal import Decimal,InvalidOperation
from zipfile import ZipFile

MAX_ROWS=65536
MAX_COLS=96


def read_biff2(source):
    """The old POS exports raw BIFF2 with 16-bit rows beyond Excel 2 limits."""
    cells={};eof=False;count=0
    with Path(source).open('rb') as stream:
        while header:=stream.read(4):
            if len(header)!=4:raise ValueError('잘린 XLS 레코드입니다.')
            kind,length=struct.unpack('<HH',header);data=stream.read(length)
            if len(data)!=length:raise ValueError('잘린 XLS 셀입니다.')
            if kind==0x0a:eof=True;break
            if kind not in {2,3,4,5,6}:continue
            if kind==6:raise ValueError('수식 셀이 있는 구형 XLS입니다. 값만 저장한 파일을 사용하세요.')
            if len(data)<7:raise ValueError('XLS 셀 구조를 확인하세요.')
            row,column=struct.unpack_from('<HH',data)
            if row>=MAX_ROWS or column>=MAX_COLS:raise ValueError('상품 목록의 행·열 허용 범위를 초과했습니다.')
            if kind==3:
                if len(data)!=15:raise ValueError('XLS 숫자 셀을 확인하세요.')
                value=struct.unpack_from('<d',data,7)[0]
                if not math.isfinite(value):raise ValueError('유한하지 않은 XLS 숫자입니다.')
                value=str(int(value)) if value.is_integer() else str(value)
            elif kind==2:
                if len(data)!=9:raise ValueError('XLS 정수 셀을 확인하세요.')
                value=str(struct.unpack_from('<H',data,7)[0])
            elif kind==4:
                if len(data)<8 or len(data)!=8+data[7]:raise ValueError('XLS 문자 셀을 확인하세요.')
                try:value=data[8:].decode('cp949')
                except UnicodeError:raise ValueError('구형 XLS 한글 인코딩을 확인하세요.') from None
            else:
                if len(data)!=9 or data[8]:raise ValueError('오류 값을 포함한 XLS입니다.')
                value='1' if data[7] else '0'
            cells.setdefault(row,{})[column]=value;count+=1
            if count>2_000_000:raise ValueError('상품 목록의 셀 수가 너무 많습니다.')
    if not eof or not cells:raise ValueError('완전한 XLS 데이터가 아닙니다.')
    width=max(max(row) for row in cells.values())+1
    if width*(max(cells)+1)>2_000_000:raise ValueError('상품 목록의 펼친 셀 수가 너무 많습니다.')
    return [[cells.get(r,{}).get(c,'') for c in range(width)] for r in range(max(cells)+1)]


def read_excel(source):
    source=Path(source)
    if not source.is_file() or source.is_symlink() or not 8<=source.stat().st_size<=32*1024*1024:raise ValueError('32MB 이하의 Excel 파일을 선택하세요.')
    with source.open('rb') as stream:header=stream.read(8)
    if source.suffix.lower()=='.xls':
        if header==b'\x09\x00\x04\x00\x02\x00\x10\x00':return read_biff2(source)
        import xlrd
        try:
            book=xlrd.open_workbook(str(source),encoding_override='cp949',on_demand=True)
            try:
                sheet=book.sheet_by_index(0)
                if sheet.nrows>MAX_ROWS or sheet.ncols>MAX_COLS or sheet.nrows*sheet.ncols>2_000_000:raise ValueError('상품 목록의 행·열 허용 범위를 초과했습니다.')
                return [[str(int(v)) if type(v) is float and v.is_integer() else str(v) for v in sheet.row_values(r)] for r in range(sheet.nrows)]
            finally:book.release_resources()
        except (xlrd.XLRDError,AssertionError,IndexError,UnicodeError) as exc:raise ValueError('XLS 파일을 읽지 못했습니다. 값만 저장한 파일을 사용하세요.') from exc
    if source.suffix.lower()!='.xlsx':raise ValueError('XLS 또는 XLSX 파일을 선택하세요.')
    from openpyxl import load_workbook
    with ZipFile(source) as archive:
        if len(archive.infolist())>2000 or sum(i.file_size for i in archive.infolist())>128*1024*1024:raise ValueError('압축 해제 크기가 너무 큰 XLSX입니다.')
    book=load_workbook(source,read_only=True,data_only=True,keep_links=False)
    try:
        sheet=book.worksheets[0]
        if sheet.max_row>MAX_ROWS or sheet.max_column>MAX_COLS or sheet.max_row*sheet.max_column>2_000_000:raise ValueError('상품 목록의 행·열 허용 범위를 초과했습니다.')
        return [[str(v) if v is not None else '' for v in row] for row in sheet.iter_rows(values_only=True)]
    finally:book.close()


def suggest_tax(name,path):
    """Only an editable proposal, never a legal determination from a name."""
    processed=('양념','조림','볶음','볶은','구운','삶은','익힌','데친','조미','튀김','스프','소스','즉석','통조림','훈제','구이','조리','맛살','어묵','소시지','햄','음료','주스','과자','라면')
    if any(word in name for word in ('양념','조미액','볶은','삶은','구운')):return 'taxable'
    if path[0] in {'야채','청과','양곡','계란','정육','수산','건어물'}:
        return 'taxable' if any(word in name for word in processed) else 'exempt'
    if path==('가공식품','가루류','밀가루'):return 'exempt'
    if path[0]=='가공식품' or path[:2] in {('비식품','생활세제류'),('비식품','세면용품'),('비식품','전기·전지'),('비식품','주방용품')}:
        return 'taxable'
    return None


def plan_excel(source):
    from .catalog_defaults import suggest_category
    rows=read_excel(source)
    header_index=next((i for i,row in enumerate(rows[:20]) if '상품명' in row and '바코드' in row),None)
    if header_index is None:raise ValueError('상품명과 바코드 제목이 있는 첫 시트를 사용하세요.')
    headers={str(value).strip():i for i,value in enumerate(rows[header_index])};result=[];skipped=0
    def value(row,*names):return next((str(row[headers[name]]).strip() for name in names if name in headers and headers[name]<len(row)), '')
    for number,row in enumerate(rows[header_index+1:],header_index+2):
        name=value(row,'상품명')
        if not name or name in {'합계','평균','표준편차'} and not value(row,'바코드','공급사'):skipped+=1;continue
        path=tuple(value(row,h) for h in ('대분류','중분류','소분류'))
        suggested=not any(path)
        if suggested:path=suggest_category(name)
        specification=value(row,'규격');unit=value(row,'기본단위','단위') or ('g' if specification.lower()=='g' else '개')
        raw_tax=value(row,'과면세','과세구분','과세/면세')
        tax={'면':'exempt','면세':'exempt','과':'taxable','과세':'taxable','exempt':'exempt','taxable':'taxable'}.get(raw_tax)
        tax_source='원본' if tax else '추천'
        if not tax:tax=suggest_tax(name,path)
        record=dict(source_row=number,name=name,barcode=value(row,'바코드'),specification=specification,supplier_name=value(row,'공급사'),category=path[0],category_mid=path[1],category_small=path[2],category_source='추천' if suggested else '원본',tax_mode=tax,tax_source=tax_source if tax else '확인 필요',base_unit=unit,price_basis=100 if unit=='g' else 1,purchase_price=value(row,'매입가','매입단가') or '0',sale_price=value(row,'판매가','판매단가') or '0',origin=value(row,'원산지'),shelf=value(row,'매대'),memo=value(row,'상품설명','메모'))
        errors=[]
        for key in ('purchase_price','sale_price'):
            try:
                price=Decimal(record[key].replace(',',''))
                if not price.is_finite() or price<0 or price*100!=int(price*100) or price>90_000_000_000:raise ValueError()
                record[key]=format(price,'f')
            except (InvalidOperation,ValueError,OverflowError):errors.append('단가 확인')
        if len(name)>300 or len(record['barcode'])>100:errors.append('상품명·바코드 길이 확인')
        if path[2] and not path[1]:errors.append('중분류 확인')
        record['error']=' / '.join(dict.fromkeys(errors));result.append(record)
    if not result:raise ValueError('등록할 상품이 없습니다.')
    return dict(rows=result,headers=list(headers),skipped=skipped,missing_tax=not any(h in headers for h in ('과면세','과세구분','과세/면세')),missing_categories=not any(h in headers for h in ('대분류','중분류','소분류')))
