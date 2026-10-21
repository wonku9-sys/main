"""Shared widgets for the independently designed offline desktop application."""
from __future__ import annotations

from datetime import date
from decimal import Decimal
from pathlib import Path
from weakref import WeakMethod

from PySide6.QtCore import QDate, QEvent, Qt, Signal, QSize
from PySide6.QtGui import QColor, QFont, QIcon, QPageLayout, QPainter, QPen, QPixmap, QTextDocument
from PySide6.QtPrintSupport import QPrintDialog, QPrinter, QPrinterInfo, QPrintPreviewDialog
from PySide6.QtWidgets import (
    QComboBox, QDateEdit, QDialog, QDialogButtonBox, QFileDialog, QFormLayout,
    QFrame, QHBoxLayout, QHeaderView, QLabel, QLineEdit, QMessageBox,
    QPushButton, QTableWidget, QTableWidgetItem, QVBoxLayout, QWidget, QToolButton, QMenu,
)
from .font_support import font_family,document_font


def won(value):
    return f"{Decimal(str(value or 0)):,.2f}"


def amount(value):
    value = Decimal(str(value or 0))
    if value.as_tuple().exponent == -2:
        return f"{value:,.2f}"
    return f"{value:,.3f}".rstrip("0").rstrip(".")


def date_field(value=None):
    field = QDateEdit()
    field.setCalendarPopup(True)
    field.setDisplayFormat("yyyy-MM-dd")
    field.setKeyboardTracking(False)
    parsed = QDate.fromString(str(value), "yyyy-MM-dd") if value else QDate.currentDate()
    field.setDate(parsed if parsed.isValid() else QDate.currentDate())
    field.setMinimumWidth(124)
    return field


def action(text, callback, primary=False, danger=False):
    result = QPushButton(text)
    from .menu_icons import menu_icon
    result.setIcon(menu_icon(text))
    result.setIconSize(QSize(20,20))
    result.setObjectName("primary" if primary else "danger" if danger else "action")
    result.clicked.connect(callback)
    mark_query_button(result,callback)
    from .menu_help import ACTION_HELP, attach_help
    attach_help(result, ACTION_HELP.get(text))
    from .permissions import restrict_button
    restrict_button(result, callback)
    return result


def action_menu(text, controls, parent):
    """Keep secondary actions available once, preserving their original permission gate."""
    from .menu_icons import menu_icon
    button=QToolButton(parent);button.setText(text);button.setIcon(menu_icon('settings'))
    button.setToolButtonStyle(Qt.ToolButtonStyle.ToolButtonTextBesideIcon)
    button.setPopupMode(QToolButton.ToolButtonPopupMode.InstantPopup)
    menu=QMenu(button);pairs=[]
    for control in controls:
        if control is None:menu.addSeparator();continue
        control.setParent(parent);control.hide()
        entry=menu.addAction(control.icon() if not control.icon().isNull() else menu_icon(control.text()),control.text())
        entry.setToolTip(control.toolTip());entry.triggered.connect(control.click);pairs.append((entry,control))
    def refresh():
        for entry,control in pairs:entry.setEnabled(control.isEnabled());entry.setText(control.text())
    menu.aboutToShow.connect(refresh);refresh();menu.setToolTipsVisible(True);button.setMenu(menu)
    button.setToolTip('보조 기능을 펼칩니다. 등록·수정·삭제 권한과 기존 확인 절차는 유지됩니다.')
    button._source_controls=controls
    return button


def reference_grid(grid, label_columns=(1,)):
    """Reference-style list surface without changing read-only flags or semantic colors."""
    grid.reference_layout=True;grid.reference_label_columns=label_columns
    grid.setShowGrid(True);grid.setAlternatingRowColors(False)
    grid.verticalHeader().setDefaultSectionSize(32)
    grid.setStyleSheet('''QTableWidget{background:#A4A4A4;gridline-color:#B9BBB7;border:1px solid #8E9589;border-radius:0;}
        QTableWidget::item{padding:2px 4px;border:0;}
        QTableWidget::item:selected{background:#DCECF9;}
        QHeaderView::section{background:#E9EACA;color:#18231C;border:0;border-right:1px solid #A4AA99;border-bottom:1px solid #A4AA99;padding:4px 3px;}''')


def output_icon(text,callback):
    button=action(text,callback)
    button.setAccessibleName(text);button.setText('');button.setFixedSize(36,32)
    button.setIconSize(QSize(28,28))
    button.setFlat(True)
    button.setStyleSheet('QPushButton{background:transparent;border:0;border-radius:0;padding:0;min-width:36px;max-width:36px;min-height:32px;max-height:32px;} QPushButton:hover,QPushButton:focus{background:palette(light);} QPushButton:pressed{background:#A8BBAE;}')
    button.setObjectName('outputPrintIcon' if text=='인쇄' else 'outputSaveIcon')
    button.setToolTip(text+' · 현재 조회 결과를 '+('미리 보고 인쇄합니다.\n단축키: Ctrl+P' if text=='인쇄' else '엑셀 파일로 저장합니다.\n단축키: Ctrl+E'))
    button.setStatusTip(button.toolTip())
    return button


def mark_query_button(button,callback):
    # Metadata must not retain a bound page/window through a Qt child cycle.
    if getattr(callback,'__name__','')=='refresh' and getattr(callback,'__self__',None) is not None:
        button._sy_query_handler=WeakMethod(callback)


def search_icon():
    from .menu_icons import menu_icon
    return menu_icon('검색')


READ_PERMISSIONS = {
    "PermissionsPage": ("admin",),
    "SupplierPaymentsPage": ("trade",), "SupplierPaymentDialog": ("trade",),
    "ProductsPage": ("catalog",), "PartnersPage": ("catalog",),
    "ProductDialog": ("catalog",), "LabelsPage": ("catalog",),
    "LabelDialog": ("catalog",), "PartnerPicker": ("catalog", "trade"),
    "UsersPage": ("admin",), "SettingsPage": ("admin",),
    "MembersPage": ("loyalty",), "MemberDialog": ("loyalty",),
    "PromotionsPage": ("loyalty",), "PromotionDialog": ("loyalty",),
    "PricesPage": ("loyalty",), "GiftsPage": ("loyalty",),
    "TradesPage": ("trade",), "LedgerPage": ("trade",), "TradeDialog": ("trade",),
    "CountsPage": ("inventory",), "CountEditor": ("inventory",),
    "InventoryPage": ("inventory",), "ReportPage": ("reports",),
    "ReportsPage": ("reports",), "OverviewPage": ("reports",),
    "CommissionPage": ("reports",), "ClosingsPage": ("inventory", "trade"),
}


def _output_context(widget, parent=None):
    """Use the nearest explicit grid/view policy, with safe class fallbacks.

    Temporary print grids may be parentless, so their supplied dialog is also
    searched. A known Store without a view policy still requires a live login.
    Unbound grids remain useful for tests and non-business utility documents.
    """
    candidates, seen = [], set()
    for start in (widget, parent):
        current = start
        while current is not None and id(current) not in seen:
            seen.add(id(current))
            candidates.append(current)
            current = current.parentWidget() if isinstance(current, QWidget) else None
    store = owner = None
    permissions = None
    for candidate in candidates:
        if store is None and getattr(candidate, "store", None) is not None:
            store = candidate.store
            if owner is None:
                owner = candidate
        if permissions is None:
            explicit = getattr(candidate, "read_permissions", None)
            if explicit is not None:
                permissions = (explicit,) if isinstance(explicit, str) else tuple(explicit)
                owner = candidate
            else:
                for cls in type(candidate).__mro__:
                    if cls.__name__ in READ_PERMISSIONS:
                        permissions, owner = READ_PERMISSIONS[cls.__name__], candidate
                        break
    return store, permissions or (), owner


def _clear_output_cache(widget, owner):
    grids = []
    for candidate in (widget, owner):
        if isinstance(candidate, DataGrid):
            grids.append(candidate)
        if isinstance(candidate, QWidget):
            grids.extend(candidate.findChildren(DataGrid))
    seen = set()
    for grid in grids:
        if id(grid) in seen:
            continue
        seen.add(id(grid))
        blocked = grid.blockSignals(True)
        try:
            grid.fill([])
        finally:
            grid.blockSignals(blocked)


def authorize_output(widget, parent=None):
    """Re-read account state immediately before exporting or painting output."""
    store, permissions, owner = _output_context(widget, parent)
    if store is None:
        return True
    try:
        store._require(*permissions)
        from .permissions import menu_key_for
        key = menu_key_for(owner) if owner is not None else None
        if key:
            store.require_menu(key)
        return True
    except PermissionError as exc:
        _clear_output_cache(widget, owner)
        target = parent if isinstance(parent, QWidget) else widget if isinstance(widget, QWidget) else None
        QMessageBox.warning(target, "출력 권한 확인", f"{exc}\n조회 자료를 지웠습니다. 다시 로그인하거나 접근 권한을 확인하세요.")
        return False


def export_rows(parent, title, headers, rows, access_widget=None):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    context = access_widget if access_widget is not None else parent
    if not authorize_output(context, parent):
        return None
    if not rows:
        QMessageBox.information(parent, "엑셀 저장", "조회된 자료가 없습니다.")
        return None
    filename, _ = QFileDialog.getSaveFileName(parent, "조회 결과 저장", f"SY_POS_{title}_{date.today().isoformat()}.xlsx", "Excel 통합문서 (*.xlsx)")
    if not filename:
        return None
    # The native file chooser runs a nested event loop. Another administrator
    # may have disabled this account while it was open.
    if not authorize_output(context, parent):
        return None
    path = Path(filename)
    if path.suffix.lower() != ".xlsx":
        path = path.with_suffix(".xlsx")
    try:
        book = Workbook()
        sheet = book.active
        sheet.title = "조회 자료"
        sheet.append(list(headers))
        money_cells=set()
        for values in rows:
            values=list(values)
            sheet.append([float(v) if isinstance(v, Decimal) else v for v in values])
            money_cells.update((sheet.max_row,index+1) for index,value in enumerate(values) if isinstance(value,Decimal) and value.as_tuple().exponent==-2)
        for row in sheet.iter_rows():
            for cell in row:
                # User-supplied text must never become a spreadsheet formula.
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                elif isinstance(cell.value, (int, float)):
                    cell.number_format = "#,##0.00" if (cell.row,cell.column) in money_cells else "#,##0.###"
                cell.alignment = Alignment(vertical="center")
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="21485F")
            cell.font = Font(color="FFFFFF", bold=True)
        for column in range(1, len(headers) + 1):
            sheet.column_dimensions[get_column_letter(column)].width = min(44, max(14, len(str(headers[column - 1])) * 2 + 6))
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        if not authorize_output(context, parent):
            return None
        book.save(path)
        return path
    except (OSError, ValueError) as exc:
        QMessageBox.warning(parent, "저장 실패", str(exc))
        return None


class Cell(QTableWidgetItem):
    def __init__(self, value):
        self.value = value
        text = f"{value:,}" if isinstance(value, int) and not isinstance(value, bool) else amount(value) if isinstance(value, Decimal) else str(value if value is not None else "")
        super().__init__(text)
        self.setData(Qt.ItemDataRole.UserRole + 1, value)
        if isinstance(value, (int, float, Decimal)):
            self.setTextAlignment(Qt.AlignmentFlag.AlignRight | Qt.AlignmentFlag.AlignVCenter)
            if value < 0:
                self.setForeground(QColor("#BA5247"))

    def __lt__(self, other):
        try:
            return self.value < other.value
        except (TypeError, AttributeError):
            return self.text() < other.text()


class HeaderSafeTable(QTableWidget):
    """Size titles from header metadata only; never scan all product rows."""
    def fit_header_text(self):
        header = self.horizontalHeader()
        for column in range(self.columnCount()):
            item = self.horizontalHeaderItem(column)
            if item and not item.toolTip():
                item.setToolTip(item.text())
            if not self.isColumnHidden(column):
                needed = header.sectionSizeHint(column)
                if header.sectionSize(column) < needed:
                    header.resizeSection(column, needed)

    def showEvent(self, event):
        super().showEvent(event)
        self.fit_header_text()

    def changeEvent(self, event):
        super().changeEvent(event)
        if event.type() in (QEvent.Type.FontChange, QEvent.Type.StyleChange) and self.isVisible():
            self.fit_header_text()


class DataGrid(HeaderSafeTable):
    def __init__(self, headers):
        super().__init__(0, len(headers))
        self.setHorizontalHeaderLabels(headers)
        self.setAlternatingRowColors(True)
        self.setSelectionBehavior(QTableWidget.SelectionBehavior.SelectRows)
        self.setSelectionMode(QTableWidget.SelectionMode.SingleSelection)
        self.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
        self.verticalHeader().setDefaultSectionSize(35)
        self.horizontalHeader().setMinimumSectionSize(55)
        self.horizontalHeader().setDefaultSectionSize(125)
        self.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Interactive)
        self.horizontalHeader().setStretchLastSection(True)
        self.setSortingEnabled(True)
        self.setMinimumHeight(150)

    def fill(self, rows, records=None):
        self.setSortingEnabled(False)
        self.clearContents()
        self.setRowCount(len(rows))
        for row_index, values in enumerate(rows):
            for column, value in enumerate(values):
                item = getattr(self,'cell_type',Cell)(value)
                background=('#EEEFCF' if column in self.reference_label_columns else '#FFFFFF') if getattr(self,'reference_layout',False) else ('#E8EFF0' if row_index%2==0 else '#DDE7E9')
                item.setBackground(QColor(background))
                item.setToolTip(str(value if value is not None else ''))
                if isinstance(value,(int,Decimal)) and not isinstance(value,bool) and value<0:
                    item.setForeground(QColor('#B42318'))
                    item.setToolTip('음수 값입니다. 부족 재고·취소·차감 내역과 거래일자를 확인하세요.')
                if column == 0 and records is not None:
                    item.setData(Qt.ItemDataRole.UserRole, records[row_index])
                self.setItem(row_index, column, item)
        self.setSortingEnabled(True)

    def selected(self):
        row = self.currentRow()
        item = self.item(row, 0) if row >= 0 else None
        return item.data(Qt.ItemDataRole.UserRole) if item else None

    def export_rows(self):
        return [[self.item(r, c).data(Qt.ItemDataRole.UserRole + 1) if self.item(r, c) else "" for c in range(self.columnCount())] for r in range(self.rowCount())]

    def export(self, parent=None, title="조회"):
        if not authorize_output(self, parent):
            return None
        return export_rows(parent or self, title, [self.horizontalHeaderItem(c).text() for c in range(self.columnCount())], self.export_rows(), access_widget=self)

    def document_html(self, title="조회", info=""):
        from html import escape
        headers = "".join(f"<th>{escape(self.horizontalHeaderItem(c).text())}</th>" for c in range(self.columnCount()))
        rows = []
        for values in self.export_rows():
            cells = []
            for value in values:
                content = won(value) if isinstance(value, int) else amount(value) if isinstance(value, Decimal) else str(value if value is not None else "")
                cells.append(f"<td>{escape(content)}</td>")
            rows.append("<tr>" + "".join(cells) + "</tr>")
        return f"<html><head><style>body{{font-size:9pt;color:#172e40}}h1{{font-size:18pt}}table{{border-collapse:collapse;width:100%}}th{{background:#e7edf2}}td,th{{border:1px solid #aebfca;padding:5px}}p{{color:#526879}}</style></head><body><h1>{escape(title)}</h1><p>SY POS 독립판 · {date.today().isoformat()} · {self.rowCount():,}건<br>{escape(info).replace(chr(10), '<br>')}</p><table cellspacing='0' cellpadding='4'><thead><tr>{headers}</tr></thead><tbody>{''.join(rows)}</tbody></table></body></html>"

    def _paint_document(self, printer, title, info, parent=None):
        if not authorize_output(self, parent):
            return False
        document = QTextDocument()
        document.setDefaultFont(document_font(9))
        document.setHtml(self.document_html(title, info))
        if not authorize_output(self, parent):
            return False
        document.print_(printer)
        return True

    def to_pdf(self, filename, title="조회", info=""):
        if not authorize_output(self):
            return None
        printer = QPrinter(QPrinter.PrinterMode.HighResolution)
        printer.setOutputFormat(QPrinter.OutputFormat.PdfFormat)
        printer.setOutputFileName(str(filename))
        if self.columnCount() > 7:
            printer.setPageOrientation(QPageLayout.Orientation.Landscape)
        if self._paint_document(printer, title, info):
            return Path(filename)
        return None

    def print_pages(self, parent=None, title="조회", info=""):
        if not authorize_output(self, parent):
            return None
        if not self.rowCount():
            QMessageBox.information(parent or self, "인쇄", "인쇄할 조회 결과가 없습니다.")
            return
        printer = QPrinter(QPrinter.PrinterMode.HighResolution)
        if self.columnCount() > 7:
            printer.setPageOrientation(QPageLayout.Orientation.Landscape)
        preview = QPrintPreviewDialog(printer, parent or self)
        preview.setWindowTitle(f"{title} · 인쇄 미리보기")
        preview.resize(1100, 780)
        def paint(target):
            if not self._paint_document(target, title, info, parent):
                preview.reject()
        preview.paintRequested.connect(paint)
        preview.exec()


class Page(QWidget):
    def __init__(self, window, title, description=""):
        super().__init__()
        self.window, self.store, self.title = window, window.store, title
        self.body = QVBoxLayout(self)
        self.body.setContentsMargins(18, 16, 18, 15)
        self.body.setSpacing(12)
        from .query_bar import QueryBar
        self.header = QueryBar(title, self.refresh, description=description)
        self.actions = self.header.actions
        self.body.addWidget(self.header)
        self.grids = []

    def changeEvent(self,event):
        from shiboken6 import isValid
        if event.type()==QEvent.Type.EnabledChange and hasattr(self,'header') and isValid(self.header) and self.header.parentWidget() is not self:
            self.header.setEnabled(self.isEnabled())
        super().changeEvent(event)

    def table(self, headers,grid_type=DataGrid):
        result = grid_type(headers)
        if not self.grids:
            self.header.add_output(output_icon("인쇄", self.print))
            self.header.add_output(output_icon("엑셀 저장", self.export))
        self.body.addWidget(result, 1)
        self.grids.append(result)
        return result

    def notify(self, text):
        if hasattr(self.window, "notify"):
            self.window.notify(text)
        else:
            QMessageBox.information(self, "SY POS 독립판", text)

    def fail(self, error):
        QMessageBox.warning(self, "처리할 수 없습니다", str(error))

    def export(self):
        if authorize_output(self) and self.grids:
            self.grids[0].export(self, self.title)

    def refresh(self):
        pass

    def print(self):
        if authorize_output(self) and self.grids:
            self.grids[0].print_pages(self, self.title)


class ScanField(QLineEdit):
    """USB HID input: a scan terminator cannot activate a dialog's Save button."""
    scanned = Signal(str)

    def __init__(self, parent=None):
        super().__init__(parent)
        self.setMaxLength(80)
        self.setPlaceholderText("바코드 스캔 후 Enter · F8")
        self.setClearButtonEnabled(True)

    def keyPressEvent(self, event):
        if event.key() in (Qt.Key.Key_Return, Qt.Key.Key_Enter, Qt.Key.Key_Tab):
            value = self.text().strip("\r\n\t ")
            if value:
                self.scanned.emit(value)
            event.accept()
            return
        super().keyPressEvent(event)

    def event(self, event):
        # QWidget normally consumes Tab for focus traversal before keyPressEvent.
        if event.type() == QEvent.Type.KeyPress and event.key() == Qt.Key.Key_Tab and self.text().strip():
            self.keyPressEvent(event)
            return True
        return super().event(event)


class FormDialog(QDialog):
    def accept(self):
        from .navigation import mark_saved
        mark_saved(self)
        super().accept()

    def reject(self):
        from .navigation import confirm_dialog_exit
        if confirm_dialog_exit(self):
            super().reject()

    def __init__(self, title, parent=None):
        super().__init__(parent)
        self.setWindowModality(Qt.WindowModality.WindowModal)
        self.setWindowTitle(f"{title} · SY POS 독립판")
        self.resize(570, 440)
        self.setMinimumWidth(460)
        self.body = QVBoxLayout(self)
        self.body.setContentsMargins(22, 20, 22, 18)
        heading = QLabel(title)
        heading.setObjectName("dialogTitle")
        self.body.addWidget(heading)
        self.form = QFormLayout()
        self.form.setVerticalSpacing(11)
        self.form.setFieldGrowthPolicy(QFormLayout.FieldGrowthPolicy.AllNonFixedFieldsGrow)
        self.body.addLayout(self.form)
        self.fields = {}

    def field(self, key, label, value="", password=False):
        result = QLineEdit(str(value if value is not None else ""))
        if password:
            result.setEchoMode(QLineEdit.EchoMode.Password)
        self.fields[key] = result
        self.form.addRow(label, result)
        return result

    def combo(self, key, label, choices, current=None):
        result = QComboBox()
        for text, value in choices:
            result.addItem(text, value)
        result.setCurrentIndex(max(0, result.findData(current)))
        self.fields[key] = result
        self.form.addRow(label, result)
        return result

    def finish(self, callback):
        self.body.addStretch()
        self.buttons = QDialogButtonBox(QDialogButtonBox.StandardButton.Save | QDialogButtonBox.StandardButton.Cancel)
        self._sy_pinned_footer = self.buttons
        self.buttons.button(QDialogButtonBox.StandardButton.Save).setText("저장")
        self.buttons.button(QDialogButtonBox.StandardButton.Save).setObjectName("primary")
        self.buttons.button(QDialogButtonBox.StandardButton.Cancel).setText("취소")
        self.buttons.accepted.connect(callback)
        self.buttons.rejected.connect(self.reject)
        self.body.addWidget(self.buttons)
        return self.buttons

    def error(self, exc):
        QMessageBox.warning(self, "입력 확인", str(exc))

    def scroll_form(self):
        """Keep long desktop editors usable on a 768px-high laptop display."""
        from PySide6.QtWidgets import QScrollArea
        self.body.removeItem(self.form);self.form.setParent(None)
        container=QWidget();container.setLayout(self.form)
        scroll=QScrollArea(self);scroll.setWidgetResizable(True);scroll.setWidget(container)
        self.body.insertWidget(1,scroll,1);return scroll
