"""Standalone desktop entry point, designed independently from the PDF workflow."""
from __future__ import annotations

import argparse
from datetime import date
import os
from pathlib import Path
import sys
import tempfile

from PySide6.QtCore import QDate, QLockFile, QSize, Qt, QTimer,QLocale
from PySide6.QtGui import QAction, QFont, QIcon, QKeySequence
from PySide6.QtWidgets import (
    QApplication, QDialog, QFrame, QGridLayout, QHBoxLayout, QLabel, QMainWindow, QMessageBox, QMenu, QPushButton,
    QScrollArea, QStackedWidget, QStyle, QToolBar, QTreeWidget, QTreeWidgetItem, QHeaderView,
    QVBoxLayout, QWidget, QSizePolicy,
)

from . import __version__
from .font_support import font_family, verify_glyphs,regular_font
from .store import Store
from .permissions import has_menu
from .navigation import NavigationGuard
from .menu_help import MENU_HELP, ACTION_HELP, MenuHelpStyle, attach_help
from .menu_shortcuts import menu_keys, help_text, bind_action, COMMON_KEYS
from .widgets import DataGrid, FormDialog, Page, ScanField, action, date_field, won
from .offline_menus import REPORT_VIEWS, SALES_VIEWS, BASIC_VIEWS, MEMBER_VIEWS, FORM_VIEWS


STYLE = """
* { outline:none; }
QWidget { color:#243A48; font-size:14px; }
QMainWindow, QDialog { background:#F4F6F8; }
QLabel { background:transparent; }
QMenuBar { background:#FFFFFF; padding:4px 10px; border-bottom:1px solid #D8E0E6; }
QMenuBar::item { padding:5px 10px; margin:1px 2px; border:1px solid #DCE5DF; border-radius:4px; }
QMenuBar::item:selected, QMenu::item:selected { background:#E5F2EC; color:#1C6654; }
QMenu { background:white; padding:5px; border:1px solid #D5DFE5; }
QMenu::item { padding:8px 22px 8px 12px; }
QToolBar { background:white; border-bottom:1px solid #D9E2E6; padding:6px 9px; spacing:4px; }
QToolBar::separator { background:#B8C9C0; width:1px; margin:7px 9px; }
QToolButton { border:1px solid #D9E4DE; border-radius:5px; padding:5px 9px; background:#F9FBFA; }
QToolButton:hover, QToolButton:checked { background:#EAF3EF; border-color:#C9DFD3; }
QToolButton:disabled { color:#87958D; background:#F3F5F3; border-color:#E3E8E4; }
#navPanel { background:#FFFFFF; border-right:1px solid #D7E0E5; }
#brand { font-size:21px; font-weight:750; color:#173D43; }
#independent { font-size:11px; color:#9C6F37; }
#navTree { background:transparent; border:0; }
#navTree::item { padding:7px 6px; border-radius:4px; }
#navTree::item:selected { background:#E6F1EB; color:#175B4C; }
#navTree::item:hover { background:#F2F6F4; }
#pageTitle, #dialogTitle { font-size:22px; font-weight:700; color:#173F49; }
#muted { color:#778996; font-size:11px; }
#notice { color:#8A652A; background:#FFF3D9; border:1px solid #EFDCB4; padding:7px 12px; }
#workspaceHeader { background:#FFFFFF; border-bottom:1px solid #DCE4E8; }
#section { background:white; border:1px solid #DBE3E8; border-radius:8px; }
#metric { background:white; border:1px solid #DAE5E2; border-radius:8px; padding:16px; }
#metricValue { color:#206B58; font-size:26px; font-weight:700; }
#homeCard { background:white; border:1px solid #DAE5E2; border-radius:8px; }
QPushButton { background:white; border:1px solid #CCD9DF; border-radius:5px; padding:7px 12px; min-height:19px; }
QPushButton:hover { background:#EDF4F1; border-color:#AFC9BD; }
QPushButton:disabled { color:#58645C; background:#E0E4DF; border-color:#BBC6BD; }
#primary { background:#256E5B; color:white; border-color:#256E5B; font-weight:600; }
#primary:hover { background:#1B5948; }
#primary:disabled { color:#58645C; background:#E0E4DF; border-color:#BBC6BD; }
#danger { color:#AD5149; background:#FFF5F2; border-color:#EACAC2; }
QLineEdit, QComboBox, QDateEdit, QSpinBox, QDoubleSpinBox, QTextEdit { background:white; border:1px solid #CFDCE2; border-radius:4px; padding:5px 7px; min-height:22px; selection-background-color:#D8ECE2; selection-color:#163F36; }
QLineEdit:focus, QComboBox:focus, QDateEdit:focus { border:1px solid #438B75; }
QLineEdit:read-only { background:#EFF3F5; color:#708592; }
QComboBox::drop-down { border:0; width:23px; }
QComboBox QAbstractItemView { background:white; selection-background-color:#DDEEE5; }
QTableWidget { background:white; alternate-background-color:#F8FAFB; gridline-color:#E3E9ED; border:1px solid #DCE4E8; selection-background-color:#DEEEE6; selection-color:#194E3E; }
QTableWidget::item { padding:5px 8px; }
QHeaderView::section { background:#EDF2F5; color:#526F7D; font-weight:600; border:0; border-bottom:1px solid #D4E0E5; border-right:1px solid #DEE7EB; padding:8px; }
QTabWidget::pane { background:white; border:1px solid #D9E3E8; border-radius:5px; }
QTabBar::tab { padding:10px 14px; color:#7E929C; background:transparent; border-bottom:2px solid transparent; }
QTabBar::tab:selected { color:#226E55; border-bottom:2px solid #2A8064; font-weight:600; }
QGroupBox { background:white; border:1px solid #D9E4E8; border-radius:6px; margin-top:14px; padding-top:15px; }
QGroupBox::title { subcontrol-origin:margin; left:12px; padding:0 5px; font-weight:600; }
QScrollArea { border:0; background:transparent; }
QScrollBar:vertical { background:#EFF3F5; width:10px; }
QScrollBar::handle:vertical { background:#BECED5; border-radius:4px; min-height:25px; }
QScrollBar::add-line:vertical, QScrollBar::sub-line:vertical { height:0; }
QStatusBar { color:#81939E; font-size:10px; }
QToolTip { color:#243A48; background:#FFFBEA; border:1px solid #D9D4B9; padding:7px; }
QMenuBar {background:#FFFFFF;padding:6px 12px;}
#navPanel {background:#FCFDFB;border-right:1px solid #DCE5DF;}
#navTree::item {padding:8px 7px;border-radius:7px;}
#navTree::item:selected {background:#E2EEE7;color:#1E5745;border-left:3px solid #B2914C;}
#pageTitle,#dialogTitle {color:#204B40;}
QHeaderView::section {background:#EAF0EB;color:#466353;border-bottom:1px solid #CEDDD3;padding:9px;}
QPushButton {border-radius:7px;}
#primary {background:#2D6651;border-color:#2D6651;}
#danger {background:#FCEFED;border:1px solid #E8BFB8;color:#A13F34;}
#notice {background:#FFF0CC;border-left:4px solid #D49B34;}
QGroupBox,#section,#metric,#homeCard {border-radius:11px;}
/* Muted workstation palette with distinct, low-glare working regions. */
QMainWindow,QDialog,QWidget#popupSurface {background:#DCE1DA;}
QMenuBar,QToolBar {background:#D0D6D0;border-bottom:1px solid #89988E;}
QMenuBar::item {background:#E0E5DF;border-color:#9EAAA1;}
QMenu,QToolTip {background:#E9EDE2;border-color:#8E9C8E;}
QToolButton,QPushButton {background:#E1E7DE;border:1px solid #8F9F93;border-radius:4px;}
QToolButton:hover,QPushButton:hover {background:palette(light);border-color:#5D7E67;}
QToolButton:checked {background:#C6D9C9;border-color:#54765D;}
#navPanel {background:#DCE3DA;border-right:1px solid #8B9B90;}
#navTree::item:selected {background:#BDD2C0;color:#173F2A;border-left:3px solid #7C804D;}
#navTree::item:hover {background:palette(light);}
#workspaceHeader {background:#D6DED4;border-bottom:1px solid #94A293;}
#section,#homeCard,#metric,QGroupBox,QTabWidget::pane {background:#DFE6DC;border-color:#9EAD9C;}
QLineEdit,QComboBox,QDateEdit,QSpinBox,QDoubleSpinBox,QTextEdit {background:#ECF0E8;border-color:#8F9F91;border-radius:3px;}
QLineEdit:focus,QComboBox:focus,QDateEdit:focus {border-color:#3F7250;}
QComboBox QAbstractItemView {background:#E9EFE4;}
QTableWidget {background:#A8AFAA;alternate-background-color:#DAE3D6;gridline-color:#B8C2B4;border:1px solid #8F9D8B;}
QHeaderView::section {background:#D4DCBC;color:#34472F;border-bottom:1px solid #8F9A7C;border-right:1px solid #A4AE92;padding:6px;}
QTabBar::tab {background:#D1DBCC;color:#4E6750;padding:7px 11px;}
QTabBar::tab:selected {background:#E3EBDE;color:#214B2C;}
QTabBar::tab:hover {background:palette(light);}
QStatusBar {background:#CDD6CB;color:#4C6250;}
/* A text-only menu strip and independent icons, not a row of cards. */
QMenuBar {padding:1px 4px;background:#DBDDD8;}
QMenuBar::item {background:transparent;border:0;border-radius:0;margin:0;padding:3px 7px;}
QMenuBar {font-size:13px;}
QMenuBar::item:selected,QMenu::item:selected {background:palette(light);color:#243A48;}
QToolBar#mainToolbar {background:#DADDD7;padding:3px 6px;spacing:3px;}
QToolBar#mainToolbar QToolButton {background:transparent;border:0;border-radius:0;padding:3px 7px;}
QToolBar#mainToolbar QToolButton:hover,QToolBar#mainToolbar QToolButton:focus {background:palette(light);border:0;}
QToolBar#mainToolbar QToolButton:pressed,QToolBar#mainToolbar QToolButton:checked {background:#B7C9BA;border:0;}
QToolBar#mainToolbar QToolButton:disabled {background:transparent;border:0;color:#7E8981;}
QToolBar#mainToolbar::separator {background:#8F9A91;width:1px;margin:5px 8px;}
QStatusBar {font-size:12px;min-height:21px;}
"""

SECTIONS = [
    ("첫화면", [("home", "첫화면", ())]),
    ("1. 기초관리", [("products", "상품장", "catalog"), *[(key,title,"admin" if key=="user_access" else "catalog") for key,title in BASIC_VIEWS.items()], ("categories", "상품분류", "catalog"), ("suppliers", "공급사관리", "catalog"), ("customers", "판매처", "catalog"), ("labels", "라벨출력", "catalog"), ("users","직원관리","admin"), ("settings","환경설정 · 백업","admin")]),
    ("2. 매입/발주관리", [("purchases", "매입전표", "trade"), ("purchase_suppliers", "공급사별 매입현황", "trade"), ("purchase_categories", "분류별 매입현황", "trade"), ("purchase_products", "단품별 매입현황", "trade"), ("purchase_documents", "전표별 매입현황", "trade"), ("purchase_statement", "매입계산서", "trade"), ("payments", "결제관리", "trade"), ("ledger", "거래처 원장", "trade")]),
    ("3. 판매관리", [*[(key,value[0],"trade") for key,value in SALES_VIEWS.items()], ("sales","판매내역 · 고객반품","trade"), ("sales_cancelled","취소내역","trade"), ("sales_held","보류내역","trade"), ("quotations","견적서","trade"), ("statements","거래명세서","trade")]),
    ("5. 마케팅관리", [("promotions","특매관리","loyalty"),("prices","가격변경예약","loyalty")]),
    ("6. 재고관리", [("inventory", "현재고-단품별", "inventory"), ("opening_stock", "기초재고 등록", "inventory"), ("stock_waste", "재고 폐기", "inventory"), ("inventory_assets", "재고자산 현황", "inventory"), ("counts", "재고실사 입력 · 적용", "inventory"), ("closings", "현금 · 월재고 마감", ("inventory", "trade")), ("inventory_categories","현재고-분류별","inventory"), ("inventory_suppliers","현재고-공급사별","inventory"), ("inventory_monthly","월말재고","inventory")]),
    ("7. 레포트관리", [("overview","업무 현황","reports"), *[(key,value[0],"reports") for key,value in REPORT_VIEWS.items()], ("reports","기간별 종합집계","reports"), ("commission","공급사 판매수수료","reports")]),
    ("8. 회원관리", [("members","회원조회/등록","loyalty"), *[(key,title,"loyalty") for key,title in MEMBER_VIEWS.items()], ("gifts","사은품조회/등록 · 지급","loyalty"), ("member_tax","(세금)계산서 참고자료","loyalty")]),
    ("0. 관리자 모드", [("permissions", "권한 설정", "admin")]),
    ("나의 기록", [("diary","보안 일기장",()),("household","가계부",()),("worklog","업무일지",()),("lucky_numbers","행운의 숫자",())]),
]
SECTIONS[1][1].append(('store_profile','점포 등록정보','admin'))
SECTIONS[3][1].append(('orders','주문서','trade'))


def menu_caption(key):
    """Stable section/item numbering, independent of an employee's visible menus."""
    overrides={'home':'9-11','product_units':'1-02','categories':'1-11','suppliers':'1-12',
               'users':'1-13','user_access':'1-14','product_changes':'1-05','retirement':'1-06',
               'weighted_products':'1-10','labels':'1-08','customers':'1-15','settings':'1-21','store_profile':'1-22',
               'payments':'2-11','ledger':'2-12','diary':'0-21','household':'0-22','worklog':'0-23','lucky_numbers':'0-24'}
    for section,entries in SECTIONS:
        for index,(page,title,_) in enumerate(entries,1):
            if page==key:
                number=overrides.get(key,f'{section[0]}-{index:02d}')
                return f'{number}. {title}'
    return key


class LoginDialog(FormDialog):
    def __init__(self, store, parent=None,network_path=None):
        self.store, self.initial = store, store.users_count() == 0
        if self.initial:
            store.bootstrap_default_admin()
        super().__init__("사용자 확인", parent)
        self.setMinimumWidth(430)
        self.resize(1000, 720)
        self.setStyleSheet('QDialog {background:#FAF8F3;} QLabel {color:#354C49;} #dialogTitle {font-size:29px;color:#214A40;font-weight:700;} QLineEdit {background:#FFFFFF;border:1px solid #D7E0D9;border-radius:9px;padding:9px 12px;min-height:25px;} QLineEdit:focus {border:1px solid #648F7E;} QPushButton {border-radius:9px;padding:10px;} #primary {background:#285F50;color:#FFFFFF;border:0;} #notice {background:#F2EBDD;color:#73654C;border:0;border-radius:9px;padding:12px;}')
        self.selected_target='home'
        self.code = self.field("code", "사용자 코드")
        self.name = self.confirm = None
        self.password = self.field("password", "비밀번호", password=True)
        default_available = store.default_login_available()
        self.new_password=self.new_confirm=None
        if default_available:
            self.new_password=self.field('new_password','처음 사용할 새 암호 (8~128자)',password=True)
            self.new_confirm=self.field('new_confirm','새 암호 확인',password=True)
        self.code.setText("admin" if default_available else "")
        if default_available:
            self.password.setText("1234")
        self.hint = QLabel("초기 아이디: admin / 초기 비밀번호: 1234\n최초 로그인 시 본인만의 새 암호를 반드시 설정하세요. 변경 후 초기 암호는 사용할 수 없습니다." if default_available else "설정된 계정과 비밀번호로 로그인하세요.\n변경한 비밀번호는 초기 비밀번호로 되돌아가지 않습니다.")
        self.hint.setWordWrap(True)
        self.hint.setObjectName("notice" if default_available else "muted")
        self.body.addWidget(self.hint)
        self.login_worker = None
        box = self.finish(self.start_submit)
        save = box.button(box.StandardButton.Save)
        save.setText("로그인")
        save.setAutoDefault(False)
        save.setDefault(False)
        self.password.returnPressed.connect(self.start_submit)
        self.body.addWidget(action("로그인 암호 찾기 · 이메일 인증",self.reset_by_email))
        self.body.addWidget(action("이메일 주소 등록 · 발송 설정",self.setup_recovery_email))
        from .lucky_audio import login_sound_settings
        self.body.addWidget(action('로그인·경고 사운드 설정',lambda:login_sound_settings(self)))
        if network_path:
            from .shared_ui import NetworkSettingsDialog
            self.body.addWidget(action('PC 연결 설정',lambda:NetworkSettingsDialog(network_path,store,self).exec()))
        from .personal_icons import personal_icon
        self.login_choices={}
        contents=QVBoxLayout()
        while self.body.count():contents.addItem(self.body.takeAt(0))
        heading=self.findChild(QLabel,'dialogTitle');heading.setText('반가워요, 어서 오세요');heading.setWordWrap(True)
        subtitle=QLabel('일과 일상을 담는 나만의 공간');subtitle.setStyleSheet('color:#81908A;font-size:13px;');contents.insertWidget(1,subtitle)
        from .branding import LOGIN_CHOICES_STYLE
        left=QFrame();left.setObjectName('loginChoices');left.setStyleSheet(LOGIN_CHOICES_STYLE)
        choices=QVBoxLayout(left);choices.setContentsMargins(20,28,20,24);choices.setSpacing(14)
        from .branding import brand_widget
        choices.addWidget(brand_widget(left,large=True))
        brand=QLabel('나의 매장, 나의 일상');brand.setStyleSheet('font-size:16px;font-weight:600;');choices.addWidget(brand)
        choices.addWidget(QLabel('오늘은 어디로 갈까요?'));left.setMinimumWidth(245)
        for key,title in [('home','매장 업무'),('household','나의 가계부'),('diary','나의 일기장')]:
            button=action(title,lambda checked=False,target=key:self.choose_target(target));button.setIcon(personal_icon(key));button.setIconSize(QSize(36,36));button.setCheckable(True);button.setAutoExclusive(True);choices.addWidget(button);self.login_choices[key]=button
        from .lucky_numbers import open_lucky_numbers
        choices.addWidget(action('행운의 숫자 · 로그인 없이',lambda:open_lucky_numbers(self)))
        choices.addStretch();guide=QLabel('업무도, 나의 작은 기록도\n편안한 한 곳에서.\n\n같은 계정으로 로그인합니다.\n개인 암호는 별도로 확인해요.');guide.setStyleSheet('color:#CADBD0;font-size:12px;');guide.setWordWrap(True);choices.addWidget(guide)
        row=QHBoxLayout();row.addWidget(left);row.addLayout(contents,1);self.body.addLayout(row)
        self.choose_target('home')

    def choose_target(self,key):
        self.selected_target=key;self.login_choices[key].setChecked(True)
        self.buttons.button(self.buttons.StandardButton.Save).setText({'home':'매장 업무 로그인','household':'가계부로 로그인','diary':'일기장으로 로그인'}[key])

    def reset_by_email(self):
        from .email_ui import EmailCodeDialog
        if EmailCodeDialog(self.store,"login_reset",self,self.code.text()).exec():
            self.password.clear();self.hint.setText("암호가 변경되었습니다. 새 비밀번호로 로그인하세요.")

    def setup_recovery_email(self):
        from .email_ui import RecoveryEmailDialog
        try:
            if self.code.text().strip()=='admin' and self.store.default_login_available():
                raise ValueError("먼저 최초 로그인에서 새 암호를 설정한 뒤, 상단 ‘복구 이메일’에서 주소와 발송 계정을 등록하세요.")
            self.store.login(self.code.text(),self.password.text())
            try:RecoveryEmailDialog(self.store,self).exec()
            finally:self.store.logout()
            self.hint.setText("이메일 설정을 확인했습니다. 업무를 시작하려면 다시 로그인하세요.")
        except (ValueError,PermissionError) as exc:self.error(exc)

    def submit(self):
        try:
            replacing=self.code.text().strip()=='admin' and self.store.default_login_available()
            if replacing:
                value=self.new_password.text() if self.new_password else ''
                if not self.new_confirm or value!=self.new_confirm.text() or not 8<=len(value)<=128:raise ValueError('처음 사용할 새 로그인 암호를 8~128자로 입력하고 확인해 주세요.')
            self.store.login(self.code.text(), self.password.text())
            if not has_menu(self.store.actor,self.selected_target):
                self.store.logout();raise PermissionError('이 메뉴의 사용 권한이 없습니다.')
            if replacing:
                self.store.change_password(self.password.text(),value)
                self.new_password.clear();self.new_confirm.clear()
            if self.store.popup_sound_enabled():
                from .lucky_audio import login_sound
                login_sound().play()
            self.accept()
        except (ValueError, PermissionError) as exc:
            self.error(exc)

    def start_submit(self):
        # The LAN client owns a separate private Store; do not move it between threads.
        if type(self.store) is not Store or self.store.path is None:
            return self.submit()
        if self.login_worker and self.login_worker.isRunning():
            return
        try:
            replacing = self.code.text().strip() == 'admin' and self.store.default_login_available()
            replacement = self.new_password.text() if replacing and self.new_password else ''
            if replacing and (not self.new_confirm or replacement != self.new_confirm.text() or not 8 <= len(replacement) <= 128):
                raise ValueError('처음 사용할 새 로그인 암호를 8~128자로 입력하고 확인해 주세요.')
            self.store.actor = None
            self.store.lock_diary(); self.store.lock_household()
            from .background_store import StoreTask
            self.login_worker = StoreTask(self, self.store, 'login',
                (self.code.text(), self.password.text(), replacement, self.selected_target), login=True)
            self.setEnabled(False)
            self.hint.setText('로그인 확인 중입니다… 잠시 기다려 주세요.')
            self._login_result = self._login_error = None
            self.login_worker.succeeded.connect(lambda result: setattr(self, '_login_result', result))
            self.login_worker.failed.connect(lambda message: setattr(self, '_login_error', message))
            self.login_worker.finished.connect(self.finish_background_login)
            self.login_worker.start()
        except (ValueError, PermissionError) as exc:
            self.error(exc)

    def finish_background_login(self):
        self.setEnabled(True)
        if self._login_error is not None:
            self.login_failed(self._login_error)
        elif self._login_result is not None:
            self.login_completed(self._login_result)
        self._login_result = None

    def login_completed(self, session):
        try:
            row = self.store.db.execute('SELECT * FROM users WHERE id=?', (session['actor']['id'],)).fetchone()
            if (not row or not row['active'] or row['password_hash'] != session['password_hash']
                    or self.store._current_restore_token() != session['restore_token']):
                raise PermissionError('로그인 중 계정 또는 장부가 변경되었습니다. 다시 로그인하세요.')
            self.store.actor = self.store._public_user(row)
            self.store._session_password_hash = session['password_hash']
            self.store._restore_token = session['restore_token']
            self.store.require_menu(self.selected_target)
            if self.new_password:
                self.new_password.clear(); self.new_confirm.clear()
            if self.store.popup_sound_enabled():
                from .lucky_audio import login_sound
                login_sound().play()
            self.accept()
        except (ValueError, PermissionError) as exc:
            self.store.actor = None
            self.login_failed(str(exc))

    def login_failed(self, message):
        self.setEnabled(True)
        self.hint.setText('로그인 정보를 확인하고 다시 시도하세요.')
        self.error(message)

    def reject(self):
        if self.login_worker and self.login_worker.isRunning():
            return
        super().reject()

    def closeEvent(self, event):
        if self.login_worker and self.login_worker.isRunning():
            event.ignore()
        else:
            event.accept(); super().reject()


from .home_ui import HomePage
class OverviewPage(Page):
    def __init__(self, window):
        super().__init__(window, "이번 달 업무 현황", "이번 달 1일~오늘 · 확정된 거래와 반품을 집계합니다.")
        self.metrics = {}
        row = QHBoxLayout()
        for key, title in (("purchase_total", "순매입"), ("sale_total", "순매출"), ("profit", "매출총이익")):
            panel = QFrame()
            panel.setObjectName("metric")
            layout = QVBoxLayout(panel)
            layout.addWidget(QLabel(title))
            number = QLabel("0 원")
            number.setObjectName("metricValue")
            layout.addWidget(number)
            self.metrics[key] = number
            row.addWidget(panel)
        self.body.addLayout(row)
        self.grid = self.table(["일자", "단위", "순매입", "순매출", "매출원가", "매출총이익"])

    def refresh(self):
        today = date.today()
        report = self.store.report(today.replace(day=1).isoformat(), today.isoformat(), "day")
        for key, widget in self.metrics.items():
            widget.setText(f"{won(report['totals'].get(key, 0))} 원")
        self.grid.fill([[r["label"], r["unit"], r["purchase_total"], r["sale_total"], r["cost"], r["profit"]] for r in report["rows"]], report["rows"])


class InventoryPage(Page):
    def __init__(self, window):
        super().__init__(window, "현재고 · 입출고", "기본단위로 환산한 현재 수량과 등록순 이동평균 원가입니다. 확장단위 스캔도 같은 재고를 사용합니다.")
        filters = QHBoxLayout()
        self.search = ScanField()
        self.search.setPlaceholderText("상품코드 · 상품명 · 규격 · 기본/확장 바코드, F8 스캔")
        self.search.scanned.connect(lambda _: self.refresh())
        filters.addWidget(self.search, 1)
        filters.addWidget(action("조회", self.refresh))
        filters.addWidget(action("입출고 이력", self.history))
        self.body.addLayout(filters)
        self.grid = self.table(["코드", "상품", "분류", "현재고", "기본단위", "최소재고", "재고원가액", "관리", "보관 위치"])
        self.total = QLabel()
        self.body.addWidget(self.total)

    def refresh(self):
        from decimal import Decimal
        rows = self.store.products(self.search.text().strip())
        self.grid.fill([[r["code"], r["name"], r["category"], Decimal(r["stock_qty"]), r["base_unit"], Decimal(r["min_qty"]), r["stock_cost"], "관리" if r["tracked"] else "제외", f"{r['warehouse']} / {r['shelf']}"] for r in rows], rows)
        self.total.setText(f"{len(rows):,}개 상품 · 재고 원가액 {won(sum(r['stock_cost'] for r in rows if r['tracked']))} 원")

    def focus_scan(self):
        self.search.setFocus()
        self.search.selectAll()

    def history(self):
        selected = self.grid.selected()
        if not selected:
            return
        from decimal import Decimal
        dialog = FormDialog(f"{selected['name']} · 입출고", self)
        dialog.resize(1000, 630)
        grid = DataGrid(["일자", "구분", "수량 증감", "잔여 수량", "원가 증감", "잔여 원가", "사유"])
        records = self.store.stock_events(selected["id"])
        grid.fill([[r["date"], r["kind"], Decimal(r.get("qty_delta", "0")), Decimal(r.get("balance_qty", "0")), r["cost_delta"], r["balance_cost"], r["reason"]] for r in records], records)
        dialog.body.addWidget(grid, 1)
        dialog.body.addWidget(action("엑셀 저장", lambda: grid.export(dialog, "입출고 이력")))
        dialog.body.addWidget(action("인쇄", lambda: grid.print_pages(dialog, "입출고 이력")))
        dialog.body.addWidget(action("닫기", dialog.reject))
        dialog.exec()


class CommissionPage(Page):
    def __init__(self, window):
        super().__init__(window, "공급사 판매 수수료", "판매 시 기록한 공급사·수수료율 기준입니다. 카드 승인·카드사 정산과 무관한 로컬 판매 수수료입니다.")
        filters = QHBoxLayout()
        self.start = date_field(date.today().replace(day=1).isoformat())
        self.end = date_field()
        filters.addWidget(self.start)
        filters.addWidget(QLabel("~"))
        filters.addWidget(self.end)
        filters.addStretch()
        filters.addWidget(action("조회", self.refresh, primary=True))
        self.body.addLayout(filters)
        self.grid = self.table(["공급사", "순매출", "판매 수수료", "차인지급 기준액"])

    def refresh(self):
        report = self.store.commission_report(self.start.date().toString("yyyy-MM-dd"), self.end.date().toString("yyyy-MM-dd"))
        records = report["rows"] if isinstance(report, dict) else report
        self.grid.fill([[r.get("supplier_name", r.get("label", "")), r.get("sale_total", r.get("net_sales", 0)), r.get("commission_fee", r.get("commission", r.get("commission_amount", 0))), r.get("payable", r.get("net_payable", 0))] for r in records], records)


class LabelsPage(Page):
    def __init__(self, window):
        super().__init__(window, "라벨 · 바코드 출력", "상품명·가격·원산지·Code128/EAN13 바코드를 라벨 프린터로 출력합니다.")
        self.body.addWidget(QLabel("USB 리더기: 키보드 입력(HID), 스캔 후 Enter. 프린터: Windows에 설치된 제조사 드라이버 사용."))
        self.body.addWidget(QLabel("기본 60×40mm, 40×30mm 등 사용자 크기·출력 위치·수량 설정. PDF·미리보기로 확인 후 실제 1장 시험 출력을 권장합니다."))
        self.body.addWidget(action("상품 · 확장단위 라벨 선택", self.open_labels, primary=True))
        self.body.addStretch()

    def open_labels(self):
        from .labels import LabelDialog
        LabelDialog(self.store, parent=self).exec()


class MainWindow(QMainWindow):
    def __init__(self, store, demo=False,network_path=None,popup_menus=False):
        super().__init__()
        self.store, self.demo = store, demo
        from .appearance import apply_theme
        self._appearance_prefs=store.appearance_preferences()
        apply_theme(self._appearance_prefs['theme'],self._appearance_prefs['holiday'])
        from .ui_sounds import SoundController
        self.sounds=SoundController(self)
        self.popup_menus=popup_menus;self.popup_windows={};self.personal_windows={}
        self.network_path=network_path;self.shared_service=None;self.shared_refresh=None
        self.pages, self.items, self.nav_actions, self.scrollers = {}, {}, {}, {}
        self.page_surfaces = {}
        from .menu_history import MenuHistory
        self.menu_history = MenuHistory()
        self._history_replaying = False
        self.history_actions = {}
        self.menu_keys = menu_keys(SECTIONS)
        self.purchase_navigation_item = None
        self.current_page, self.current_key = None, None
        self._store_name=store.store_profile().get('store_name','') or '내 점포'
        self._title_page_key=None
        self._backup_runtime_enabled, self._exit_backup_done = False, False
        self._exit_confirmed=False
        self._payment_auth_scope=None
        from .branding import app_icon, brand_widget
        self.setWindowTitle(f"SY POS 독립판 {__version__} · 단독 사용 / 선택형 PC 공유" + (" [체험]" if demo else ""))
        self.setWindowIcon(app_icon())
        self.resize(1500, 930)
        self.setMinimumSize(920, 600)
        self.statusBar().setSizeGripEnabled(True)
        central = QWidget()
        layout = QHBoxLayout(central)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.setSpacing(0)
        self.setCentralWidget(central)
        sidebar = QFrame()
        sidebar.setObjectName("navPanel")
        sidebar.setFixedWidth(215)
        side = QVBoxLayout(sidebar)
        side.setContentsMargins(0, 0, 0, 7);side.setSpacing(0)
        from .query_bar import QueryContext, QUERY_BAR_HEIGHT
        self.sidebar_heading=QFrame();self.sidebar_heading.setObjectName('sidebarHeading')
        self.sidebar_heading.setFixedHeight(QUERY_BAR_HEIGHT)
        self.sidebar_heading.setStyleSheet('QFrame#sidebarHeading{background:#C6CAD0;border-bottom:1px solid #87939C;}')
        heading_row=QHBoxLayout(self.sidebar_heading);heading_row.setContentsMargins(3,4,3,4)
        self.selected_menu=QueryContext('');self.selected_menu.setObjectName('selectedMenuBadge')
        self.selected_menu.setProperty('menuBadge',True);self.selected_menu.setStyleSheet('font-weight:600;font-size:14px;color:#253B58;')
        heading_row.addWidget(self.selected_menu);side.addWidget(self.sidebar_heading)
        self.tree = QTreeWidget()
        self.tree.setObjectName("navTree")
        # Keep the reference's compact hierarchy with comfortable 14px text
        # instead of shrinking the menu to fit. Long labels stay scrollable.
        self.tree.setStyleSheet('QTreeWidget#navTree{font-size:14px;} QTreeWidget#navTree::item{padding:3px 5px;min-height:19px;}')
        self.tree.setHeaderHidden(True)
        self.tree.setIconSize(QSize(18,18))
        self.tree.setIndentation(12)
        self.tree.setTextElideMode(Qt.TextElideMode.ElideNone)
        self.tree.header().setStretchLastSection(False)
        self.tree.header().setSectionResizeMode(QHeaderView.ResizeMode.ResizeToContents)
        self.tree.setHorizontalScrollMode(QTreeWidget.ScrollMode.ScrollPerPixel)
        self.tree.setExpandsOnDoubleClick(False)
        self.tree.itemClicked.connect(self.sidebar_clicked)
        self.tree.itemExpanded.connect(self.sidebar_expanded)
        side.addWidget(self.tree, 1)
        self.personal_panel=QFrame()
        self.personal_panel.setStyleSheet("QFrame {background:transparent;border:0;} QLabel {background:transparent;border:0;color:#526A74;font-weight:600;} QPushButton {text-align:left;background:transparent;border:1px solid transparent;padding:10px;} QPushButton:hover,QPushButton:focus {background:palette(light);border-color:#91A7B2;}")
        self.personal_layout=QHBoxLayout(self.personal_panel);self.personal_layout.setContentsMargins(12,7,12,7);self.personal_layout.setSpacing(9)
        self.personal_layout.addWidget(QLabel("나의 기록"))
        self.personal_navigation_root=QTreeWidgetItem(["나의 기록"])
        side.addWidget(QLabel("● 로컬 SQLite · 서버 불필요"))
        side.addWidget(QLabel("자료 저장: data-v2\n기존 프로그램 DB와 별도"))
        layout.addWidget(sidebar)
        right = QVBoxLayout()
        right.setSpacing(0)
        header = QFrame()
        header.setObjectName("workspaceHeader")
        top = QHBoxLayout(header)
        top.setContentsMargins(20, 8, 20, 8)
        # Keep the programmatic status target, not a repeated visible heading.
        self.location = QLabel(header)
        self.location.hide()
        top.addStretch()
        top.addWidget(QLabel(f"{store.actor['name']} · {'관리자 모드' if store.actor['role'] == 'admin' else '일반 사용자'}"))
        self.admin_mode_button = None
        if store.actor["role"] == "admin":
            self.admin_mode_button = action("관리자 설정", lambda: self.open_page("settings"))
            top.addWidget(self.admin_mode_button)
        top.addWidget(action("비밀번호 변경", self.password))
        top.addWidget(action("복구 이메일",self.recovery_email))
        top.addWidget(action('오늘의 글귀',self.show_inspiration))
        if self.network_path:top.addWidget(action('PC 연결 · 공유',self.configure_network))
        top.addWidget(action("로그아웃", self.logout))
        # Account controls must not force a wide minimum window size.
        header_scroll=QScrollArea(self);header_scroll.setObjectName('accountHeaderScroll')
        header_scroll.setWidgetResizable(True);header_scroll.setFrameShape(QFrame.Shape.NoFrame)
        header_scroll.setVerticalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOff)
        header_scroll.setWidget(header)
        header_scroll.setFixedHeight(header.sizeHint().height()+self.style().pixelMetric(QStyle.PixelMetric.PM_ScrollBarExtent))
        # Preserve the account handlers, exposed below through 0. 기타 rather
        # than inserting another persistent row between icons and the query.
        header_scroll.hide()
        self.account_controls=header
        self.shell_notices=[]
        if demo:
            banner = QLabel("체험 모드 · 실제 거래자료와 분리된 샘플입니다. 스캔·반품·실사·라벨을 시험할 수 있습니다.")
            banner.setObjectName("notice")
            self.shell_notices.append(banner)
        elif store.actor["code"] == "admin" and store.default_login_available():
            banner = QLabel("초기 비밀번호를 사용 중입니다. 상단 ‘비밀번호 변경’에서 새 비밀번호로 변경해 주세요.")
            banner.setObjectName("notice")
            self.shell_notices.append(banner)
        self.shared_notice=QLabel();self.shared_notice.setObjectName('notice');self.shared_notice.setWordWrap(True);self.shared_notice.hide()
        from .query_bar import NavigationBar
        self.navigation_bar=NavigationBar('첫화면',self)
        self.navigation_bar.setFixedHeight(0);self.navigation_bar.hide()
        self.stack = QStackedWidget()
        right.addWidget(self.stack, 1)
        for banner in self.shell_notices:right.addWidget(banner)
        right.addWidget(self.shared_notice)
        self.personal_dock=QFrame();self.personal_dock.setObjectName('personalDock');self.personal_dock.setStyleSheet('QFrame#personalDock{background:palette(window);border-top:1px solid #B7C5CA;}')
        dock=QHBoxLayout(self.personal_dock);dock.setContentsMargins(12,7,12,7);dock.addStretch();dock.addWidget(self.personal_panel);dock.addStretch();right.addWidget(self.personal_dock)
        layout.addLayout(right, 1)
        self.toolbar = QToolBar("업무 바로가기", self)
        self.toolbar.setObjectName('mainToolbar')
        self.toolbar.setMovable(False)
        self.toolbar.setToolButtonStyle(Qt.ToolButtonStyle.ToolButtonTextUnderIcon)
        self.toolbar.setIconSize(QSize(38, 38))
        self.addToolBar(self.toolbar)
        # Line 1 is the native title bar: small SY app icon, page, store, version.
        # Line 2 contains only numbered textual business menu groups.
        self.utility_menu=QMenu('0. 기타',self.menuBar())
        for title, entries in SECTIONS:
            allowed = [entry for entry in entries if self.allowed(entry[2]) and has_menu(self.store.actor, entry[0])]
            if not allowed:
                continue
            if title=='첫화면':menu=self.utility_menu
            elif title=='나의 기록':menu=self.utility_menu.addMenu(title)
            elif title=='0. 관리자 모드':
                menu=self.utility_menu;self.menuBar().addMenu(menu)
            else:menu = self.menuBar().addMenu(title)
            menu.setToolTipsVisible(True)
            attach_help(menu.menuAction(), " · ".join(caption for _, caption, _ in allowed) + " 업무 메뉴를 엽니다.")
            # Home remains available through the top icon and keyboard/history,
            # but must not duplicate the entry in the left business menu.
            parent = (QTreeWidgetItem([title]) if title=="첫화면" else
                      self.personal_navigation_root if title=="나의 기록" else QTreeWidgetItem(self.tree, [title]))
            if title == "2. 매입/발주관리":
                self.purchase_navigation_item = parent
            parent.setExpanded(False)
            font = parent.font(0)
            font.setBold(True)
            parent.setFont(0, font)
            for key, caption, _permission in allowed:
                item = parent if key == "home" else QTreeWidgetItem(parent, [caption])
                item.setData(0, Qt.ItemDataRole.UserRole, key)
                hint = help_text(MENU_HELP[key], self.menu_keys[key])
                item.setToolTip(0, hint)
                item.setStatusTip(0, hint)
                if title!='나의 기록':item.setIcon(0,self.style().standardIcon(QStyle.StandardPixmap.SP_DirHomeIcon if key=='home' else QStyle.StandardPixmap.SP_FileIcon))
                self.items[key] = item
                from .menu_icons import menu_icon
                item.setIcon(0,menu_icon(key))
                nav = QAction(caption, self)
                nav.setIcon(menu_icon(key));nav.setIconVisibleInMenu(True)
                attach_help(nav, hint)
                nav.setCheckable(True)
                nav.setShortcuts([QKeySequence(sequence) for sequence in self.menu_keys[key]])
                nav.setShortcutContext(Qt.ShortcutContext.WidgetWithChildrenShortcut)
                nav.setAutoRepeat(False)
                nav.triggered.connect(lambda checked=False, page=key: self.go_home() if page=='home' else self.open_page(page))
                menu.addAction(nav)
                self.addAction(nav)
                self.nav_actions[key] = nav
                if title=="나의 기록":
                    from .personal_icons import personal_icon
                    button=action({"diary":"일기장","household":"가계부","worklog":"업무일지","lucky_numbers":"행운의 숫자"}[key],lambda checked=False,target=key:self.open_page(target))
                    button.setAccessibleName(caption)
                    button.setIcon(personal_icon(key));button.setIconSize(QSize(26,26))
                    button.ensurePolished()
                    button.setMinimumWidth(button.fontMetrics().horizontalAdvance(button.text())+button.iconSize().width()+36)
                    attach_help(button,hint);self.personal_layout.addWidget(button)
        self.personal_panel.setVisible(bool(self.personal_navigation_root.childCount()))
        self.personal_dock.setVisible(bool(self.personal_navigation_root.childCount()))
        # 0. 기타 must remain reachable for employees without admin permissions.
        if self.utility_menu.menuAction() not in self.menuBar().actions():self.menuBar().addMenu(self.utility_menu)
        self.utility_menu.addSeparator();self.account_actions={}
        for button in self.account_controls.findChildren(QPushButton):
            tool=self.utility_menu.addAction(button.text());tool.triggered.connect(button.click)
            attach_help(tool,button.toolTip() or f'{button.text()} 기능을 실행합니다.')
            self.account_actions[button.text()]=tool
        from .menu_icons import menu_icon
        for title, callback in (('이전', self.previous_menu), ('다음', self.next_menu)):
            tool = bind_action(self, title, COMMON_KEYS[title], callback)
            tool.setIcon(menu_icon(title))
            self.toolbar.addAction(tool)
            self.history_actions[title] = tool
        self.toolbar.addSeparator()
        for title, callback in (("첫화면", self.go_home), ('달력',self.open_calendar), ("상품장", lambda: self.open_page("products")), ("매입전표", lambda: self.open_page("purchases")), ("판매조회", lambda: self.open_page("sales")), ("F8 스캔", self.focus_scan)):
            if title == "상품장" and not self.allowed("catalog") or title in ("매입전표", "판매조회") and not self.allowed("trade"):
                continue
            if title=='F8 스캔':self.toolbar.addSeparator()
            from .menu_icons import menu_icon
            tool = QAction(menu_icon(title), title, self)
            attach_help(tool, ACTION_HELP[title])
            tool.triggered.connect(callback)
            self.toolbar.addAction(tool)
        self.toolbar_brand_space=QWidget(self.toolbar)
        self.toolbar_brand_space.setObjectName('toolbarBrandSpace')
        self.toolbar_brand_space.setSizePolicy(QSizePolicy.Policy.Expanding,QSizePolicy.Policy.Preferred)
        self.toolbar_brand_space.setStyleSheet('QWidget#toolbarBrandSpace{background:transparent;border:0;}')
        brand_row=QHBoxLayout(self.toolbar_brand_space);brand_row.setContentsMargins(0,0,0,0);brand_row.setSpacing(0)
        self.toolbar_brand=brand_widget(self.toolbar_brand_space)
        brand_row.addStretch(1);brand_row.addWidget(self.toolbar_brand);brand_row.addStretch(1)
        self.toolbar_brand_action=self.toolbar.addWidget(self.toolbar_brand_space)
        self.toolbar_brand_action.setToolTip(self.toolbar_brand.toolTip())
        self.common_shortcut_actions = {
            title: bind_action(self, title, COMMON_KEYS[title], callback)
            for title, callback in (("F8 스캔", self.focus_scan), ("조회", self.refresh),
                                    ("엑셀 저장", self.export), ("인쇄", self.print),
                                    ("새 항목", self.new_current), ('달력',self.open_calendar), ("단축키 안내", self.show_shortcut_guide))
        }
        guide = self.common_shortcut_actions['단축키 안내']
        attach_help(guide, ACTION_HELP['단축키 안내'])
        self.utility_menu.addSeparator();self.utility_menu.addAction(guide)
        from .retired_games import offer_cleanup
        cleanup=self.utility_menu.addAction('등록 게임 파일 삭제');cleanup.triggered.connect(lambda:offer_cleanup(self))
        attach_help(cleanup,'미니 오락실에 복사 등록했던 게임 사본만 확인 후 삭제합니다. 원본·외부 실행기·점수·업무자료는 보존합니다.')
        # The guide remains in 0. 기타/F1; no extra text-only toolbar button.
        self.notify(f"새 독립판 DB: {store.path} · F8 바코드  F5 조회  Ctrl+E 엑셀  Ctrl+P 인쇄")
        self.navigation = NavigationGuard(self)
        if self.items:
            self.open_page("home")
        self.maintenance_timer = QTimer(self)
        self.maintenance_timer.setInterval(60 * 60 * 1000)
        self.maintenance_timer.timeout.connect(self.maintain_backups)
        self.popup_timer=QTimer(self);self.popup_timer.setInterval(3000);self.popup_timer.timeout.connect(self.refresh_popups)
        if self.popup_menus:self.popup_timer.start()
        self.backup_timer = QTimer(self)
        self.backup_timer.setInterval(60_000)
        self.backup_timer.timeout.connect(self.scheduled_backup)
        self.alert_timer = QTimer(self)
        self.alert_timer.setInterval(30_000)
        self.alert_timer.timeout.connect(self.show_desk_alerts)
        from .ambient_effects import AmbientEffects
        self.ambient_effects=AmbientEffects(self,self._appearance_prefs)

    def update_appearance(self):
        self._appearance_prefs=self.store.appearance_preferences()
        self.ambient_effects.configure(self._appearance_prefs)

    def enable_runtime(self):
        self.navigation.enable()
        from .unsaved_work import CrashJournal
        self.crash_journal=CrashJournal(self)
        recover=QAction('미저장 작업 복구',self);recover.triggered.connect(self.crash_journal.show_center);self.utility_menu.addAction(recover)
        recover.setToolTip('비정상 종료 전에 암호화 보관된 임시 작업을 확인하고 입력창으로 복구합니다. 복구만으로 저장·재고 확정은 하지 않습니다.')
        if self.crash_journal.rows():QTimer.singleShot(900,self.crash_journal.offer)
        self._backup_runtime_enabled = True
        self.alert_timer.start()
        QTimer.singleShot(2000,self.show_desk_alerts)
        self.pages["home"].enable_online()

    def show_desk_alerts(self):
        if not self._backup_runtime_enabled or not self.isVisible() or QApplication.activeModalWidget() is not None:
            return
        try:
            records = self.store.due_desk_alerts()
            if records:
                from .alert_ui import ScheduleAlertDialog
                ScheduleAlertDialog(records,self).exec()
                if self.current_key=="home": self.current_page.refresh_events()
        except (ValueError,PermissionError) as exc:self.notify(f"일정 알림 확인 실패: {exc}")

    def show_inspiration(self,new=False):
        from .inspiration import InspirationDialog
        try:
            dialog=InspirationDialog(self.store,new,self);dialog.exec();dialog.deleteLater()
        except (ValueError,PermissionError,OSError) as exc:self.notify(str(exc))

    def scheduled_backup(self):
        # Automatic backups run only during normal exit, never while entering data.
        return None

    def prepare_exit(self):
        from .independent_windows import close_independent_windows
        if not close_independent_windows(self):return False
        if not self.navigation.allow_leave(self.current_page):
            return False
        if self.shared_service:self.shared_service.begin_close()
        if self._backup_runtime_enabled and not self._exit_backup_done and self.store.actor:
            try:
                self.store.backup_on_exit()
                self._exit_backup_done = True
            except (ValueError, PermissionError, OSError) as exc:
                answer = QMessageBox.question(self, "종료 백업 실패", f"종료 백업을 저장하지 못했습니다.\n{exc}\n\n백업 없이 종료할까요?", QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No, QMessageBox.StandardButton.No)
                if answer != QMessageBox.StandardButton.Yes:
                    if self.shared_service:self.shared_service.cancel_close()
                    return False
                self._exit_backup_done = True
        return True

    def refresh_popups(self):
        if not self.popup_windows or QApplication.activeModalWidget() is not None:return
        visible=[popup for popup in self.popup_windows.values() if popup.isVisible()]
        if not visible:return
        try:revision=self.store.audit_revision()
        except (ValueError,PermissionError,OSError):return
        for popup in visible:
            if popup.revision==revision or self.navigation.changed(popup.page):continue
            try:
                self.store.require_menu(popup.key);popup.page.refresh();self.navigation.track(popup.page);popup.revision=revision
            except (ValueError,PermissionError,OSError):
                for grid in popup.page.findChildren(DataGrid):grid.fill([])
                popup.page.setEnabled(False)

    def maintain_backups(self):
        if not self.isVisible() or QApplication.activeModalWidget() is not None:
            return
        try:
            result = self.store.run_backup_retention()
            if result["deleted"] or result["failed"]:
                message = f"오래된 백업 {len(result['deleted']):,}개를 삭제했습니다.\n최근 생성된 백업 15개는 그대로 보관합니다."
                if result["failed"]:
                    message += f"\n사용 중인 파일 {len(result['failed']):,}개는 삭제하지 못했습니다."
                QMessageBox.information(self, "주간 백업 정리", message)
                if "settings" in self.pages and self.store.actor and self.store.actor["role"] == "admin":
                    self.pages["settings"].refresh_backups()
        except (ValueError, PermissionError, OSError):
            self.notify("백업 정리를 완료하지 못했습니다. 설정에서 백업 폴더를 확인하세요.")

    def confirm_exit(self):
        message = ('SY POS를 종료하시겠습니까?\n\n'
                   '저장하지 않은 내용이 있으면 먼저 확인합니다.\n'
                   '종료 백업이 설정되어 있으면 백업을 마친 뒤 안전하게 닫습니다.\n'
                   '계속 사용하려면 아니요를 눌러 주세요.')
        return QMessageBox.question(self,'SY POS 종료',message,QMessageBox.StandardButton.Yes|QMessageBox.StandardButton.No,QMessageBox.StandardButton.No)==QMessageBox.StandardButton.Yes

    def closeEvent(self, event):
        from shiboken6 import isValid
        if any(isValid(worker) and worker.isRunning() for worker in getattr(QApplication.instance(),'_sy_catalog_workers',[])):
            QMessageBox.warning(self,'작업 처리 중','분석·등록 작업이 완료된 후 프로그램을 종료하세요.');event.ignore();return
        if not self.navigation.suspended and not self._exit_confirmed and (self._backup_runtime_enabled or event.spontaneous()):
            if not self.confirm_exit():event.ignore();return
        if not self.prepare_exit():
            event.ignore()
            return
        self._exit_confirmed=True
        self.ambient_effects.stop()
        self.maintenance_timer.stop()
        self.backup_timer.stop()
        self.alert_timer.stop()
        self.popup_timer.stop()
        self.navigation.disable()
        if hasattr(self,'crash_journal'):self.crash_journal.stop()
        self.sounds.stop()
        if self.shared_refresh:self.shared_refresh.stop()
        if self.shared_service:self.shared_service.stop()
        self.store.lock_diary();self.store.lock_household()
        super().closeEvent(event)

    def go_home(self):
        return self.navigation.request_home()

    def open_calendar(self):
        """Open the existing live shop calendar, without another popup or model."""
        if not self.go_home():return False
        self.current_page.calendar.setFocus()
        self.scrollers['home'].ensureWidgetVisible(self.current_page.calendar)
        self.notify(ACTION_HELP['달력'])
        return True

    def show_shortcut_guide(self, parent=None):
        from .menu_shortcuts import ShortcutGuide
        guide = getattr(self, '_shortcut_guide', None)
        if guide is None:
            guide = self._shortcut_guide = ShortcutGuide(self, parent if isinstance(parent, QWidget) else self)
        guide.showNormal(); guide.raise_(); guide.activateWindow()

    def recovery_email(self):
        from .email_ui import RecoveryEmailDialog
        RecoveryEmailDialog(self.store,self).exec()

    def configure_network(self):
        from .shared_ui import NetworkSettingsDialog
        dialog=NetworkSettingsDialog(self.network_path,self.store,self)
        try:dialog.exec()
        finally:
            dialog.code.clear();self.navigation.snapshots.pop(dialog,None);dialog.deleteLater()

    def reset_navigation(self):
        self.tree.collapseAll()

    def sidebar_clicked(self,item,column=0):
        key=item.data(0,Qt.ItemDataRole.UserRole)
        if key:self.open_page(key)
        elif item.parent() is None and item.childCount():item.setExpanded(True)

    def sidebar_expanded(self,item):
        if item.parent() is not None:return
        self._last_sidebar_group=item
        # Do not collapse branches inside Qt's own expand traversal. In
        # particular expandAll() may still be iterating those native items.
        QTimer.singleShot(0,self,self._collapse_other_sidebar_groups)

    def _collapse_other_sidebar_groups(self):
        item=getattr(self,'_last_sidebar_group',None)
        if item is None or not item.isExpanded():return
        for index in range(self.tree.topLevelItemCount()):
            other=self.tree.topLevelItem(index)
            if other is not item:other.setExpanded(False)

    def allowed(self, permission):
        actor = self.store.actor
        permissions = (permission,) if isinstance(permission, str) else permission
        if not permissions:
            return bool(actor)
        return bool(actor and (actor["role"] == "admin" or set(permissions) & set(actor["permissions"])))

    def open_page(self, key):
        if key not in self.items:
            return False
        if key=='payments' and self._backup_runtime_enabled and not self.verify_payment_entry():return False
        if key=='lucky_numbers':
            try:self.store.require_menu(key)
            except (ValueError,PermissionError) as exc:QMessageBox.warning(self,'행운의 숫자',str(exc));return False
            from .lucky_numbers import open_lucky_numbers
            open_lucky_numbers(self);self._record_menu_visit(key);return True
        if not (self.popup_menus and key!='home') and getattr(self, "navigation", None) and not self.navigation.allow_leave(self.current_page):
            return False
        from .catalog_ui import ProductsPage, CategoriesPage, PartnersPage, UsersPage, SettingsPage
        from .workflow_ui import TradesPage, LedgerPage, CountsPage, ReportPage, ClosingsPage
        from .loyalty_ui import MembersPage, PromotionsPage, PricesPage, GiftsPage
        from .payment_ui import SupplierPaymentsPage
        from .permission_ui import PermissionsPage
        from .supplier_ui import SuppliersPage
        factories = {"overview": lambda: OverviewPage(self), "products": lambda: ProductsPage(self), "suppliers": lambda: PartnersPage(self, "supplier"), "customers": lambda: PartnersPage(self, "customer"), "purchases": lambda: TradesPage(self, "purchase"), "sales": lambda: TradesPage(self, "sale"), "ledger": lambda: LedgerPage(self), "inventory": lambda: InventoryPage(self), "counts": lambda: CountsPage(self), "closings": lambda: ClosingsPage(self), "members": lambda: MembersPage(self), "promotions": lambda: PromotionsPage(self), "prices": lambda: PricesPage(self), "gifts": lambda: GiftsPage(self), "reports": lambda: ReportPage(self), "commission": lambda: CommissionPage(self), "labels": lambda: LabelsPage(self), "users": lambda: UsersPage(self), "settings": lambda: SettingsPage(self)}
        factories["categories"] = lambda: CategoriesPage(self)
        from .store_profile_ui import StoreProfilePage
        factories['store_profile']=lambda:StoreProfilePage(self)
        factories["payments"] = lambda: SupplierPaymentsPage(self)
        factories["permissions"] = lambda: PermissionsPage(self)
        factories["home"] = lambda: HomePage(self)
        from .stock_management_ui import StockAdjustmentPage, InventoryAssetsPage
        factories['opening_stock']=lambda:StockAdjustmentPage(self,'opening')
        factories['stock_waste']=lambda:StockAdjustmentPage(self,'waste')
        factories['inventory_assets']=lambda:InventoryAssetsPage(self)
        from .inventory_status_ui import InventoryStatusPage
        for page_key,group in (('inventory','products'),('inventory_categories','categories'),('inventory_suppliers','suppliers'),('inventory_monthly','monthly')):
            factories[page_key]=lambda group=group:InventoryStatusPage(self,group)
        factories["suppliers"] = lambda: SuppliersPage(self)
        from .diary_ui import DiaryPage,WorkLogPage
        factories["diary"]=lambda:DiaryPage(self)
        factories["worklog"]=lambda:WorkLogPage(self)
        from .household_ui import HouseholdPage
        factories["household"]=lambda:HouseholdPage(self)
        from .offline_ui import SalesHistoryPage
        factories['sales_cancelled']=lambda:SalesHistoryPage(self,'cancelled')
        factories['sales_held']=lambda:SalesHistoryPage(self,'draft')
        from .purchase_reports import PURCHASE_VIEWS, PurchaseStatusPage
        for group, (page_key, _) in PURCHASE_VIEWS.items():
            factories[page_key] = lambda group=group: PurchaseStatusPage(self, group)
        from .offline_ui import OfflineReportPage, BasicRecordsPage, MemberViewPage, SalesFormsPage
        for page_key in {**REPORT_VIEWS, **SALES_VIEWS}:
            factories[page_key] = lambda key=page_key: OfflineReportPage(self,key)
        for page_key in BASIC_VIEWS:
            factories[page_key] = lambda key=page_key: BasicRecordsPage(self,key)
        for page_key in MEMBER_VIEWS:
            factories[page_key] = lambda key=page_key: MemberViewPage(self,key)
        for page_key in FORM_VIEWS:
            factories[page_key] = lambda key=page_key: SalesFormsPage(self,key)
        try:
            self.store.require_menu(key)
            if key not in self.pages:
                page = factories[key]()
                page.menu_key = key
                permissions = next(permission for _, entries in SECTIONS for page_key, _, permission in entries if page_key == key)
                page.read_permissions = (permissions,) if isinstance(permissions, str) else permissions
                scroll = QScrollArea()
                scroll.setWidgetResizable(True)
                scroll.setWidget(page)
                self.pages[key], self.scrollers[key] = page, scroll
                from .query_bar import PageSurface
                surface=PageSurface(page,scroll,lambda checked=False,key=key:self.refresh_page(key))
                self.page_surfaces[key]=surface
                self.stack.addWidget(surface)
            page = self.pages[key]
            if key in self.popup_windows and self.popup_windows[key].isVisible():
                self.popup_windows[key].showNormal();self.popup_windows[key].raise_();self.popup_windows[key].activateWindow()
                self._sync_menu_context(key,page)
                return True
            page.setEnabled(True)
            page.refresh()
            self.navigation.track(page)
        except (ValueError, PermissionError) as exc:
            if isinstance(exc, PermissionError) and key in self.pages:
                for grid in self.pages[key].findChildren(DataGrid):
                    grid.fill([])
                self.pages[key].setEnabled(False)
            QMessageBox.warning(self, "조회 실패", str(exc))
            return False
        self.current_key, self.current_page = key, page
        self.query_bar=page.header
        self.sounds.play('menu')
        if self.popup_menus and key in {'diary','household'}:
            # Keep the settings/landing page cached, but do not open a second
            # shell window behind the actual private editor.
            (page.open_diary if key=='diary' else page.open_book)()
        elif self.popup_menus and key!='home' or key in self.popup_windows:
            from .independent_windows import PageWindow
            if key not in self.popup_windows:
                self.stack.removeWidget(self.page_surfaces[key])
                self.popup_windows[key]=PageWindow(self,key,page,self.page_surfaces[key])
            popup=self.popup_windows[key];popup.show();popup.raise_();popup.activateWindow()
        else:self.stack.setCurrentWidget(self.page_surfaces[key])
        self._sync_menu_context(key,page)
        return True

    def verify_payment_entry(self):
        try:
            self.store.require_menu('payments')
            scope=(self.store.actor['id'],getattr(self.store,'_restore_token',None),getattr(self.store,'restore_epoch',None))
            if self._payment_auth_scope==scope:return True
            from .payment_access_ui import PaymentAccessDialog
            dialog=PaymentAccessDialog(self)
            if dialog.exec()!=QDialog.DialogCode.Accepted or not dialog.verified:return False
            self._payment_auth_scope=scope;return True
        except (ValueError,PermissionError,OSError) as exc:
            QMessageBox.warning(self,'결제관리 확인',str(exc));return False

    def _sync_menu_context(self,key,page,record=True):
        self.current_key,self.current_page=key,page
        payment_mode=key=='payments'
        self.toolbar.setVisible(not payment_mode);self.menuBar().setVisible(not payment_mode)
        self.tree.parentWidget().setVisible(not payment_mode)
        self.personal_dock.setVisible(not payment_mode and bool(self.personal_navigation_root.childCount()))
        self.query_bar=page.header
        caption=menu_caption(key)
        if payment_mode:
            from .payment_workspace import view_caption
            caption=view_caption(page)
        self.selected_menu.setText(caption);self.selected_menu.setStatusTip(caption)
        page.header.context.setText(caption)
        page.header.context.setVisible(payment_mode or bool(self.popup_menus and key in self.popup_windows))
        if payment_mode:
            page.header.context.setText(caption)
        page.header.description.hide()
        if key in {'settings','store_profile'} or self._title_page_key in {'settings','store_profile'}:
            self._store_name=self.store.store_profile().get('store_name','') or '내 점포'
        self._title_page_key=key
        self.setWindowTitle(f'{caption} ◀ {" ".join(self._store_name.split())} ◀ SY POS {__version__}'+(' [체험]' if self.demo else ''))
        self.navigation_bar.context.setText(page.title)
        # Menu execution changes the workspace, not the sidebar arrangement.
        vertical=self.tree.verticalScrollBar().value();horizontal=self.tree.horizontalScrollBar().value()
        expanded=[(self.tree.topLevelItem(index),self.tree.topLevelItem(index).isExpanded()) for index in range(self.tree.topLevelItemCount())]
        self.tree.setCurrentItem(self.items[key] if self.items[key].treeWidget() is self.tree else None)
        for item,was_expanded in expanded:item.setExpanded(was_expanded)
        self.tree.verticalScrollBar().setValue(vertical);self.tree.horizontalScrollBar().setValue(horizontal)
        self.location.setText(f"매장 업무  /  {page.title}")
        for page_key, nav in self.nav_actions.items():
            nav.setChecked(page_key == key)
        self.notify(page.header.description.text() or MENU_HELP.get(key,caption))
        if record:self._record_menu_visit(key)
        else:self._update_history_actions()

    def _record_menu_visit(self,key):
        if not self._history_replaying:self.menu_history.record(key)
        self._update_history_actions()

    def _update_history_actions(self):
        for title,step in (('이전',-1),('다음',1)):
            key=self.menu_history.label(step)
            caption=self.items[key].text(0) if key in self.items else ''
            hint=help_text(ACTION_HELP[title],COMMON_KEYS[title])
            if caption:hint+='\n이동할 메뉴: '+caption
            actions=[self.history_actions.get(title)]
            actions.extend(getattr(popup,'common_shortcut_actions',{}).get(title) for popup in self.popup_windows.values())
            for tool in actions:
                if tool is not None:
                    tool.setEnabled(key is not None)
                    tool.setToolTip(hint);tool.setStatusTip(hint)

        page=self.pages.get('payments')
        if page is not None:
            from .payment_workspace import update_history
            update_history(page)

    def previous_menu(self):
        page=self.pages.get('payments')
        if self.current_key=='payments' and page is not None and page.payment_history_index>0:
            from .payment_workspace import travel
            travel(page,-1);return True
        return self._travel_menu_history(-1)

    def next_menu(self):
        page=self.pages.get('payments')
        if self.current_key=='payments' and page is not None and page.payment_history_index<len(page.payment_history)-1:
            from .payment_workspace import travel
            travel(page,1);return True
        return self._travel_menu_history(1)

    def _travel_menu_history(self,step):
        if self._history_replaying:return False
        target=self.menu_history.target(step)
        if target is None:return False
        self._history_replaying=True
        try:
            success=self.open_page(self.menu_history.entries[target])
        finally:self._history_replaying=False
        if success:self.menu_history.index=target
        self._update_history_actions()
        return success

    def notify(self, text):
        self.statusBar().showMessage(str(text))

    def cached_products(self):
        """Shared, revision-checked product list for dialogs (e.g. purchase/
        sale entry) that need the full catalog but should not re-read every
        row from SQLite each time one is opened."""
        self.store._require('catalog', 'trade', 'inventory', 'reports', 'loyalty')
        key = (str(self.store.actor), self.store.catalog_revision())
        if getattr(self, '_products_cache_key', None) != key:
            products = self.store.products()
            self._products_cache_key = key
            self._products_cache = products
        return self._products_cache

    def refresh(self):
        self.refresh_page(self.current_key)

    def refresh_page(self,key):
        # A visible background menu must query its own page even while an
        # independent window has keyboard focus/current_page ownership.
        page=self.pages.get(key)
        if page is None:return
        self._refresh_page(page,key)

    def _refresh_page(self,page,key):
        if page:
            try:
                if not self.navigation.allow_leave(page):
                    return
                self.store.require_menu(key)
                page.refresh()
                self.navigation.track(page)
            except (ValueError, PermissionError) as exc:
                if isinstance(exc, PermissionError):
                    for grid in page.findChildren(DataGrid):
                        grid.fill([])
                    page.setEnabled(False)
                QMessageBox.warning(self, "조회", str(exc))

    def new_trade(self, kind):
        if not self.allowed("trade"):
            return
        if self.open_page("purchases" if kind == "purchase" else "sales"):
            self.current_page.new_trade()

    def new_current(self):
        if self.current_key in ("purchases", "sales"):
            self.new_trade("purchase" if self.current_key == "purchases" else "sale")
        elif self.current_page and hasattr(self.current_page, "add"):
            self.current_page.add()

    def focus_scan(self):
        modal = QApplication.activeModalWidget()
        target = modal or self.current_page
        if target and hasattr(target, "focus_scan"):
            target.focus_scan()
        elif target:
            field = target.findChild(ScanField)
            if field:
                field.setFocus()
                field.selectAll()

    def export(self):
        if self.current_page:
            self.current_page.export()

    def print(self):
        if self.current_page:
            self.current_page.print()

    def password(self):
        dialog = FormDialog("내 비밀번호 변경", self)
        old = dialog.field("old", "현재 비밀번호", password=True)
        new = dialog.field("new", "새 비밀번호", password=True)
        confirmation = dialog.field("confirm", "새 비밀번호 확인", password=True)
        def save():
            try:
                if new.text() != confirmation.text():
                    raise ValueError("새 비밀번호 확인이 일치하지 않습니다.")
                self.store.change_password(old.text(), new.text())
                dialog.accept()
            except (ValueError, PermissionError) as exc:
                dialog.error(exc)
        dialog.finish(save)
        dialog.exec()

    def logout(self):
        if QMessageBox.question(self, "로그아웃", "로그아웃하고 프로그램을 종료할까요?") == QMessageBox.StandardButton.Yes:
            if self.prepare_exit():
                self.store.logout()
                self.navigation.suspended = True
                self.close()

    def finish_restore(self, result):
        self.navigation.suspended = True
        self._exit_backup_done = True
        self.maintenance_timer.stop()
        self.setEnabled(False)
        for page in self.pages.values():
            page.hide()
        title='업무자료 초기화 완료' if result.get('business_reset') else '백업 복원 완료'
        message=('업무자료를 초기화했습니다. 개인 기록과 기존 로그인 계정·암호는 보존했습니다.\n프로그램을 종료합니다. 다시 실행한 뒤 기존 계정으로 로그인하세요.\n\n초기화 전 보호 백업:\n' if result.get('business_reset') else '백업 자료를 불러왔습니다. 프로그램을 종료합니다.\n다시 실행한 뒤 백업 당시 계정·비밀번호로 로그인하세요.\n\n복원 전 자료의 자동 백업:\n')
        QMessageBox.information(self,title,message+result['safety_backup'])
        self.close()
        QApplication.instance().quit()


def data_directory():
    if getattr(sys, "frozen", False):
        return Path(sys.executable).resolve().parent / "data-v2"
    root = Path(os.environ.get("LOCALAPPDATA", os.environ.get("XDG_DATA_HOME", str(Path.home() / ".local" / "share"))))
    return root / "SeoyoungPOS-Standalone" / "data-v2"


def main(argv=None):
    parser = argparse.ArgumentParser(description="SY POS 독립판 · 서버 없는 오프라인 매장 업무")
    parser.add_argument("--data-dir", type=Path)
    parser.add_argument("--demo", action="store_true")
    parser.add_argument("--smoke-test", action="store_true")
    parser.add_argument("--screenshot", type=Path)
    args = parser.parse_args(argv)
    app = QApplication.instance() or QApplication(sys.argv[:1])
    from .popup_policy import install_popup_policy
    install_popup_policy(app)
    QLocale.setDefault(QLocale('ko_KR'))
    from .branding import app_icon
    app.setApplicationName("SY POS 독립판")
    app.setWindowIcon(app_icon())
    app.setOrganizationName("SeoyoungStandalone")
    app.setStyle(MenuHelpStyle())
    try:
        app.setFont(regular_font(10))
    except RuntimeError as exc:
        if args.smoke_test:
            raise
        QMessageBox.critical(None, "글꼴 파일 오류", str(exc))
        return 1
    from .branding import spin_button_style
    app.setStyleSheet(STYLE+spin_button_style())
    from .lucky_audio import prewarm_mouse_sounds
    QTimer.singleShot(0,app,prewarm_mouse_sounds)
    scratch = tempfile.TemporaryDirectory(prefix="seoyoung-v2-") if args.smoke_test else None
    directory = Path(scratch.name) if scratch else args.data_dir or data_directory()
    is_demo = args.demo or args.smoke_test
    connection_path=directory/'pc-connection.json';network_config={'mode':'standalone'}
    path = directory / ("seoyoung-standalone-demo.db" if is_demo else "seoyoung-standalone-v2.db")
    store, backdrop, window, process_lock = None, None, None, None
    try:
        directory.mkdir(parents=True, exist_ok=True)
        if not is_demo:
            from .shared_config import load_config
            try:network_config=load_config(connection_path)
            except ValueError as exc:
                QMessageBox.warning(None,'PC 연결 설정',str(exc))
                from .shared_ui import NetworkSettingsDialog
                NetworkSettingsDialog(connection_path).exec();return 0
            if network_config['mode']=='client':
                from hashlib import sha256
                path=directory/'personal'/sha256(network_config['certificate'].encode()).hexdigest()[:24]/'personal.db'
                path.parent.mkdir(parents=True,exist_ok=True)
        process_lock = QLockFile(str(path) + ".lock")
        process_lock.setStaleLockTime(0)
        if not process_lock.tryLock(0):
            raise ValueError("이 장부를 사용하는 프로그램이 이미 실행 중입니다. 기존 창을 종료한 뒤 다시 실행하세요.")
        if network_config['mode']=='client':
            from .shared_client import ClientStore
            try:store=ClientStore(network_config,path)
            except (ValueError,OSError) as exc:
                QMessageBox.warning(None,'기준 PC에 연결할 수 없습니다',str(exc))
                from .shared_ui import NetworkSettingsDialog
                NetworkSettingsDialog(connection_path).exec();return 0
        else:store = Store(path)
        login_target='home'
        if is_demo:
            from .demo import seed
            seed(store)
        else:
            backdrop = QMainWindow()
            backdrop.setWindowTitle("SY POS 독립판 · 사용자 확인")
            backdrop.resize(1400, 900)
            for section, _ in SECTIONS:
                backdrop.menuBar().addMenu(section).setEnabled(False)
            backdrop.show()
            login=LoginDialog(store,backdrop,connection_path)
            if login.exec() != QDialog.DialogCode.Accepted:
                return 0
            login_target=login.selected_target;login.password.clear()
        window = MainWindow(store, is_demo,None if is_demo else connection_path,popup_menus=False)
        if network_config['mode']=='host':
            from .shared_server import SharedService
            service=SharedService(path,network_config)
            try:service.start();window.shared_service=service;store._shared_host_running=True
            except (OSError,ValueError) as exc:
                QMessageBox.warning(window,'장부 공유 시작 실패',f'{exc}\n\n이 PC 단독 업무는 계속 사용할 수 있습니다. PC 연결 설정과 주소·포트를 확인하세요.')
        if network_config['mode'] in {'host','client'}:
            from .shared_ui import SharedRefresh
            window.shared_refresh=SharedRefresh(window)
            window.setWindowTitle(window.windowTitle()+(' · 기준 PC' if network_config['mode']=='host' else f" · 연결 PC {network_config['address']}"))
        if not args.smoke_test:
            window.enable_runtime()
        if app.platformName() != "offscreen" and app.primaryScreen():
            available = app.primaryScreen().availableGeometry()
            window.resize(min(1500, available.width() - 20), min(930, available.height() - 35))
        window.show()
        if login_target in {'household','diary'}:
            def open_personal():
                if not args.smoke_test:window.show_inspiration(True)
                if window.open_page(login_target):
                    if not window.popup_menus:
                        page=window.current_page
                        (page.open_book if login_target=='household' else page.open_diary)()
            QTimer.singleShot(0,open_personal)
        elif not args.smoke_test:QTimer.singleShot(0,lambda:window.show_inspiration(True))
        if backdrop:
            backdrop.close()
        if args.smoke_test or args.screenshot:
            def verify():
                try:
                    if args.smoke_test:
                        verify_glyphs()
                        from .menu_icons import ATLAS,ATLAS_SHA256,_atlas_image
                        from hashlib import sha256
                        assert sha256(ATLAS.read_bytes()).hexdigest()==ATLAS_SHA256 and _atlas_image() is not None
                        print('PASS icon3d: original transparent 30-icon atlas, high-DPI menu artwork and offline integrity',flush=True)
                        from .menu_icons import OUTPUT_ATLAS,OUTPUT_ATLAS_SHA256,_output_atlas_image
                        assert sha256(OUTPUT_ATLAS.read_bytes()).hexdigest()==OUTPUT_ATLAS_SHA256 and _output_atlas_image() is not None
                        print('PASS realistic printer/Excel: dedicated transparent artwork, complete aspect-preserving source regions and original output commands',flush=True)
                        from .menu_icons import PERSONAL_ASSETS,PERSONAL_SHA256,_personal_image
                        assert all(sha256(path.read_bytes()).hexdigest()==PERSONAL_SHA256[kind] and _personal_image(kind) is not None for kind,path in PERSONAL_ASSETS.items())
                        print('PASS personal 3D: dedicated diary and household artwork, transparent shared icons and private window heading decoration',flush=True)
                        from .branding import wordmark_image
                        assert not wordmark_image().isNull() and wordmark_image().hasAlphaChannel()
                        tools=window.toolbar.actions()
                        scan=next(tool for tool in tools if tool.text()=='F8 스캔')
                        assert tools[tools.index(scan)+1] is window.toolbar_brand_action
                        print('PASS 3D SY POS: offline transparent wordmark immediately after F8 scan; shared login artwork',flush=True)
                        from .audio_smoke import verify_default_audio
                        decoded=verify_default_audio()
                        assert len(decoded)==6 and all(frames>0 for frames in decoded.values())
                        print('PASS default MP3: six original files decoded to PCM; no speaker playback claimed',flush=True)
                        from .demo import DEMO_CODE, DEMO_PASSWORD
                        from time import monotonic, sleep
                        async_login=LoginDialog(store,window);async_login.code.setText(DEMO_CODE);async_login.password.setText(DEMO_PASSWORD)
                        login_errors=[];async_login.error=lambda error:login_errors.append(str(error))
                        async_login.start_submit()
                        deadline=monotonic()+15
                        while (async_login.login_worker.isRunning() or store.actor is None) and monotonic()<deadline:
                            app.processEvents();sleep(.01)
                        app.processEvents()
                        assert not async_login.login_worker.isRunning() and not login_errors and store.actor, f'Background login failed: {login_errors}'
                        assert async_login.result()==QDialog.DialogCode.Accepted
                        async_login.deleteLater()
                        print('PASS background login: separate SQLite connection, UI completion and valid session',flush=True)
                        if sys.platform=='win32':
                            from .shared_smoke import run as shared_smoke
                            assert shared_smoke(directory/'isolated-network-proof')
                        assert len(store.products()) == 6
                        assert len(store.documents(status="posted")) >= 9
                        for key in window.items:
                            assert window.open_page(key), f"Desktop page unavailable: {key}"
                        assert not any(key.startswith("seoyoung_pos") for key in sys.modules), "Legacy program imported"
                        from .demo import DEMO_CODE, DEMO_PASSWORD
                        reopened = Store(path)
                        try:
                            reopened.login(DEMO_CODE, DEMO_PASSWORD)
                            assert len(reopened.products()) == 6
                            assert len(reopened.documents(status="posted")) >= 9
                        finally:
                            reopened.close()
                        backup_file = store.create_backup()
                        store.save_member({"code": "SMOKE-RESTORE", "name": "복원 검증용 임시 회원"})
                        restored = store.restore_backup(backup_file, DEMO_PASSWORD)
                        assert Path(restored["safety_backup"]).is_file()
                        store.login(DEMO_CODE, DEMO_PASSWORD)
                        assert not store.members("SMOKE-RESTORE")
                    window.open_page("products")
                    app.processEvents()
                    if args.screenshot:
                        args.screenshot.parent.mkdir(parents=True, exist_ok=True)
                        assert window.grab().save(str(args.screenshot))
                        if args.smoke_test:
                            from .workflow_ui import ProductPicker, TradeDialog, PartnerPicker
                            def capture(suffix, widget):
                                widget.show()
                                app.processEvents()
                                bar=getattr(widget,'query_bar',None)
                                if bar is not None:
                                    center=bar.query_button.mapTo(bar,bar.query_button.rect().center()).x()
                                    assert abs(center-bar.rect().center().x())<=1,'Query button must stay centered'
                                    nav=getattr(widget,'navigation_bar',None)
                                    if nav is not None and widget.isAncestorOf(bar):
                                        from PySide6.QtCore import QPoint
                                        assert bar.mapTo(widget,QPoint()).y()>=nav.mapTo(widget,QPoint(0,nav.height())).y(),'Query belongs below navigation, in the page heading'
                                target = args.screenshot.with_name(args.screenshot.stem + suffix + args.screenshot.suffix)
                                assert widget.grab().save(str(target))
                            window.open_page("home")
                            capture("-home", window)
                            history_revision=store.audit_revision()
                            window.open_page('products');window.open_page('purchases')
                            visits=list(window.menu_history.entries)
                            assert window.previous_menu() and window.current_key=='products'
                            assert window.previous_menu() and window.current_key=='home'
                            assert window.next_menu() and window.current_key=='products'
                            assert window.next_menu() and window.current_key=='purchases'
                            assert window.menu_history.entries==visits and store.audit_revision()==history_revision
                            capture('-menu-history',window)
                            assert not window.popup_menus and not window.popup_windows
                            window.sidebar_clicked(window.items['purchases'].parent())
                            capture('-purchase-workspace',window)
                            supplier_picker=PartnerPicker(window,'supplier')
                            supplier_picker.search.setText(store.trade_partners('supplier')[0]['code'])
                            capture('-purchase-supplier-picker',supplier_picker)
                            supplier_picker.choose()
                            assert supplier_picker.result()==QDialog.DialogCode.Accepted and supplier_picker.selected_record
                            assert not supplier_picker.isVisible()
                            assert not any(tool.text() in {'엑셀 저장','인쇄'} for tool in window.toolbar.actions())
                            print('PASS fixed workspace: inline purchase list, supplier search selection, output icons only',flush=True)
                            window.open_page('home')
                            print('PASS menu history: back/forward replay, fixed grouped toolbar, no business writes',flush=True)
                            purchase = TradeDialog(window, "purchase")
                            first = store.products()[0]
                            purchase.partner.setCurrentIndex(purchase.partner.findData(first["supplier_id"]))
                            for product in store.products(supplier_id=first["supplier_id"]):
                                purchase.scan_barcode(product["barcode"])
                            purchase.paid.setText("5000")
                            assert "미지급잔액" in purchase.balance_label.text()
                            capture("-purchase", purchase)
                            from .purchase_output import purchase_record,purchase_document,purchase_workbook
                            voucher=purchase_record(purchase)
                            from PySide6.QtWidgets import QTextBrowser
                            voucher_dialog=QDialog(window);voucher_dialog.setWindowTitle('매입전표 양식 미리보기');voucher_dialog.resize(850,800)
                            voucher_body=QVBoxLayout(voucher_dialog);voucher_view=QTextBrowser();voucher_doc=purchase_document(voucher);voucher_doc.setParent(voucher_view);voucher_view.setDocument(voucher_doc);voucher_body.addWidget(voucher_view)
                            capture('-purchase-paper',voucher_dialog);voucher_dialog.close()
                            output_xlsx=args.screenshot.parent/'purchase-voucher-smoke.xlsx';purchase_workbook(voucher).save(output_xlsx)
                            from openpyxl import load_workbook
                            output_book=load_workbook(output_xlsx);assert output_book.active['A1'].value=='매 입 전 표';assert output_book.active['B5'].value==voucher['supplier']['name'];output_book.close()
                            from .purchase_import import plan_files
                            from .purchase_import_ui import PurchaseImportDialog
                            imported_plan=plan_files([output_xlsx],store.purchase_import_catalog(),purchase.partner.currentData(),purchase.date.date().toString('yyyy-MM-dd'))
                            assert not imported_plan[0]['error'], imported_plan[0]['error']
                            preview_import=PurchaseImportDialog(store,window);preview_import.loaded(imported_plan);preview_import.busy(False)
                            capture('-purchase-excel-import',preview_import)
                            assert len(imported_plan[0]['lines'])==len(voucher['lines'])
                            import_result=store.import_purchase_drafts(imported_plan)
                            assert len(import_result['created'])==1
                            from .recovery import SCHEMA_VERSION
                            import_backup=store.create_backup();assert store.inspect_backup(import_backup)['schema_version']==SCHEMA_VERSION
                            assert store.import_purchase_drafts(imported_plan)['existing']==import_result['created']
                            store.delete_draft(import_result['created'][0])
                            print('PASS purchase Excel: actual workbook to draft, validated backup and duplicate identity',flush=True)
                            gram_draft=store.save_draft(dict(kind='purchase',partner_id=first['supplier_id'],date=purchase.date.date().toString('yyyy-MM-dd'),lines=[dict(product_id=first['id'],qty='125',unit_price='1000',gram_input=True)]))
                            gram_editor=TradeDialog(window,'purchase',store.document(gram_draft))
                            assert gram_editor.line_rows[0]['unit'].currentData()=='grams'
                            assert gram_editor.line_rows[0]['total'].text()=='1,250.00'
                            store.save_draft(gram_editor.payload(),gram_draft)
                            assert store.document(gram_draft)['total']==1250
                            gram_editor.close();store.delete_draft(gram_draft)
                            print('PASS gram draft: reopening and saving retain g/100g, not kg quantities',flush=True)
                            preview_import.close();app.processEvents();assert not preview_import.isVisible()
                            from .catalog_import_ui import CatalogImportDialog
                            idle_import=CatalogImportDialog(store,window);idle_import.show();app.processEvents();idle_import.close();app.processEvents();assert not idle_import.isVisible(),'Idle Excel importer must close'
                            idle_import.deleteLater()
                            picker = ProductPicker(purchase)
                            picker.set_apply_callback(purchase.apply_picked_products)
                            # Purchase entry keeps the picker open by default.
                            assert picker.keep_open.isChecked()
                            assert all(button.icon().isNull() for button in picker.text_only_buttons)
                            from .number_edit import GroupedNumberEdit
                            assert isinstance(purchase.line_rows[0]['qty'],GroupedNumberEdit)
                            assert any(',' in row['price'].text() for row in purchase.line_rows if row['product'].currentData())
                            assert picker.supplier_only.isChecked()
                            assert 0 < picker.grid.rowCount() < len(store.products())
                            picker.check_all(True)
                            capture("-product-search", picker)
                            picker.apply_checked();assert picker.isVisible() and not picker.checked and picker.applied_batches==1
                            print('PASS purchase picker continuity: open by default after apply, text-only footer, grouped numeric editors',flush=True)
                            capture('-product-search-continuous',picker)
                            picker.close()
                            purchase.close()
                            window.open_page('suppliers')
                            supplier_page=window.pages['suppliers'];supplier_page.grid.selectAll()
                            supplier_menu=supplier_page.make_context_menu()
                            assert any(a.text()=='선택 공급사 무효' for a in supplier_menu.actions())
                            capture('-supplier-bulk-selected',window.popup_windows.get('suppliers',window))
                            supplier_menu.popup(supplier_page.grid.viewport().mapToGlobal(supplier_page.grid.viewport().rect().center()))
                            app.processEvents();capture('-supplier-bulk-menu',supplier_menu)
                            supplier_menu.close()
                            source = next((store.document(row["id"]) for row in store.documents(kind="purchase", status="posted") if any(float(line["remaining_qty"]) > 0 for line in store.document(row["id"])["lines"])), None)
                            assert source is not None
                            returned = TradeDialog(window, "purchase_return", source=source)
                            assert returned.return_mode.isChecked() and returned.line_rows
                            returned.line_rows[0]["qty"].setText("0.1")
                            capture("-return", returned)
                            returned.close()
                            window.open_page("categories")
                            capture("-categories", window)
                            from .catalog_ui import ProductDialog,CategorySearchDialog
                            new_product=ProductDialog(store,parent=window);capture('-new-product-categories',new_product)
                            category_search=CategorySearchDialog(store,new_product);category_search.filters[2].setText('신상품')
                            assert category_search.grid.rowCount()>0
                            capture('-product-category-search',category_search);category_search.close();new_product.close()
                            from .menu_shortcuts import ShortcutGuide
                            shortcut_guide=ShortcutGuide(window);capture('-shortcut-guide',shortcut_guide);shortcut_guide.close()
                            for key,suffix in (('inventory_categories','-inventory-categories'),('inventory_suppliers','-inventory-suppliers'),('inventory_monthly','-inventory-monthly')):
                                assert window.open_page(key)
                                if key=='inventory_categories':window.current_page.children.setChecked(True)
                                capture(suffix,window)
                            from .supplier_ui import SupplierGroupsDialog
                            window.open_page("suppliers")
                            capture("-suppliers",window)
                            groups=SupplierGroupsDialog(window)
                            capture("-supplier-groups",groups);groups.close()
                            window.open_page("product_units")
                            capture("-product-units",window)
                            window.open_page("report_day_category")
                            capture("-reports",window)
                            from .offline_ui import SalesFormDialog
                            paper=SalesFormDialog(window,"quotations")
                            if paper.rows:
                                paper.rows[0]["qty"].setText("1")
                            capture("-quotation",paper);paper.close()
                            from .diary_ui import DiaryDialog
                            store.setup_diary("Smoke-Diary-Only-2026!","Smoke-Diary-Only-2026!")
                            diary_key=store.unlock_diary("Smoke-Diary-Only-2026!")
                            store.save_diary_music_preferences(False,6)
                            store.save_popup_sound_enabled(False)
                            diary=DiaryDialog(store,diary_key,window)
                            diary.title.setText("오늘의 작은 기록")
                            diary.editor.setPlainText("이것은 실제 개인 자료가 아닌 화면 검증용 예시입니다.\n오늘의 업무를 차근차근 마치고 잠깐 쉬어갑니다.")
                            assert diary.save()
                            capture("-diary",diary);diary.close()
                            assert not any(diary_key)
                            from .household_ui import HouseholdDialog
                            store.setup_household('Smoke-Cashbook-2026!','Smoke-Cashbook-2026!')
                            cash_key=store.unlock_household('Smoke-Cashbook-2026!')
                            today=date.today().isoformat()
                            store.save_household_entry(cash_key,{'date':today,'kind':'income','title':'생활비 예시','category':'생활','amount':'250000.00'})
                            store.save_household_entry(cash_key,{'date':today,'kind':'expense','title':'장보기 예시','category':'식비','amount':'25300.50','memo':'화면 검증용 가상 자료'})
                            store.save_household_memo(cash_key,today,'이번 달에는 작은 소비부터 차근차근 기록해 봐요.')
                            account=store.save_household_account(cash_key,{'name':'생활비 통장 (예시)','bank':'가상 은행','opening_date':today,'opening_balance':'1000000.55'})
                            car=store.save_household_vehicle(cash_key,{'name':'가족 자동차 (예시)','plate':'00가0000','model':'가상 차량','next_service_date':today,'next_service_km':'30000'})
                            receipt=store.save_household_entry(cash_key,{'date':today,'kind':'expense','title':'주유 (예시)','category':'차량','amount':'65000.25','account_id':account,'vehicle_id':car,'vehicle_type':'fuel','liters':'35.5','odometer':'30000','vehicle_vendor':'가상 주유소','fuel_kind':'휘발유','full_tank':True})
                            store.save_household_fixed_cost(cash_key,{'title':'통신비 (예시)','amount':'35000.55','monthly_day':str(date.today().day),'start_date':today,'account_id':account})
                            store.save_household_budget(cash_key,today[:7],'300000.00')
                            store.save_household_plan(cash_key,dict(kind='asset',name='예·적금 예시',amount='3500000.25',asset_type='deposit',date=today))
                            store.save_household_plan(cash_key,dict(kind='debt',name='남은 대출 예시',amount='1250000.50',date=today))
                            store.save_household_plan(cash_key,dict(kind='goal',name='여행 자금 예시',amount='2000000.00',account_id=account,target_date=f'{date.today().year+1}-12-31'))
                            store.save_household_plan(cash_key,dict(kind='category_budget',name='식비',amount='200000.00',month=today[:7]))
                            store.save_household_plan(cash_key,dict(kind='forecast',name='예상 월 수입',amount='2500000.00'))
                            store.due_household_fixed_alerts(cash_key);store.due_household_vehicle_alerts(cash_key)
                            store.add_personal_attachment(cash_key,'household','entry',receipt,'화면검사.txt',b'Fictional receipt for screenshot validation only.')
                            cashbook=HouseholdDialog(store,cash_key,window)
                            # The isolated capture deliberately creates due
                            # vehicle reminders below. Do not open their modal
                            # user notifications inside a headless EXE check.
                            # Production timers and reminder tests are unchanged.
                            cashbook.timer.stop();capture('-household',cashbook)
                            for index,suffix in [(1,'-accounts'),(3,'-fixed-costs'),(4,'-vehicles')]:cashbook.finances.setCurrentIndex(index);capture(suffix,cashbook)
                            cashbook.finances.setCurrentIndex(5)
                            for index,suffix in [(0,'-financial-overview'),(1,'-savings-goals'),(2,'-category-budgets'),(3,'-annual-household'),(6,'-custom-household-reports')]:cashbook.finances.planning.setCurrentIndex(index);capture(suffix,cashbook)
                            store.save_work_log(dict(date=today,title='물류 작업·인수인계 (예시)',body='실제 자료가 아닌 화면 검증용입니다.',details=dict(
                                status='done',category='입출고',customer='가상 고객사',handover='다음 근무 시 보관 구역 확인',
                                handling_revenue='250000.50',rent_cost='20000.25',capacity='100',occupied='85',
                                workers=[dict(name='가상 작업자',regular_hours='8',overtime_hours='1.25',hourly_rate='12000.50',overtime_factor='1.5')],
                                movements=[dict(activity='하역',unit='BOX',qty='500',hours='2'),dict(activity='적치',unit='PLT',qty='13',hours='1.25')]
                            )))
                            window.open_page('worklog');capture('-worklog-dashboard',window)
                            window.current_page.grid.selectRow(0);window.current_page.selection_changed();window.current_page.tabs.setCurrentIndex(2);capture('-worklog-productivity',window)
                            window.open_page('sales_realtime');capture('-sales-comparison',window)
                            from .attachment_ui import AttachmentDialog
                            attachment=AttachmentDialog(store,cash_key,'household','entry',receipt,cashbook);capture('-attachments',attachment);attachment.reject()
                            from .personal_schedule_ui import PersonalScheduleDialog
                            personal_schedule=PersonalScheduleDialog(store,cashbook);capture('-personal-schedule',personal_schedule);personal_schedule.reject()
                            from .inspiration import InspirationDialog
                            inspiration=InspirationDialog(store,True,window);capture('-inspiration',inspiration)
                            first_quote=inspiration.phrase.text();inspiration.next_button.click();inspiration.previous_button.click()
                            assert inspiration.phrase.text()==first_quote and inspiration.isVisible()
                            inspiration.explanation_button.click();assert inspiration.explanation.isVisible()
                            capture('-inspiration-explanation',inspiration);inspiration.close_button.click()
                            from .shared_ui import NetworkSettingsDialog
                            connection=NetworkSettingsDialog(directory/'smoke-connection.json',store,window);capture('-pc-connection',connection);connection.close()
                            store.save_health_entry('household',cash_key,{'category':'growth','date':today,'title':'성장 기록 예시','subject':'아이 (예시)','height':'96.7','weight':'14.23','birth_date':f'{date.today().year-3}-01-01','notes':'실제 개인 정보가 아닌 화면 검증용 가상 자료'})
                            store.save_health_entry('household',cash_key,{'category':'weight','date':today,'title':'체중 예시','subject':'나 (예시)','height':'170','weight':'68.25'})
                            store.save_health_entry('household',cash_key,{'category':'goal','date':today,'title':'개인 목표 예시','subject':'나 (예시)','target_weight':'67','target_water_ml':'1800'})
                            store.save_health_entry('household',cash_key,{'category':'meal','date':today,'title':'점심 예시','subject':'나 (예시)','meal_slot':'lunch','calories':'650.25','protein_g':'25.5','carbs_g':'70.2','fat_g':'20','food_items':'직접 입력한 예시 식단'})
                            store.save_health_entry('household',cash_key,{'category':'water','date':today,'title':'물 한 잔 예시','subject':'나 (예시)','water_ml':'250'})
                            from .health_ui import HealthDialog
                            health=HealthDialog(store,'household',cash_key,cashbook);capture('-health',health);health.reject()
                            from .health_report_ui import HealthReportDialog
                            health_report=HealthReportDialog(store,'household',cash_key,cashbook);capture('-health-report',health_report);health_report.reject()
                            from .vehicle_management_ui import VehicleReportDialog,VehicleObligationsDialog
                            store.save_household_vehicle_obligation(cash_key,{'vehicle_id':car,'title':'보험 갱신 (예시)','obligation_kind':'insurance','due_date':today,'amount':'150000.55'})
                            vehicle_report=VehicleReportDialog(store,cash_key,car,cashbook);capture('-vehicle-report',vehicle_report);vehicle_report.reject()
                            obligations=VehicleObligationsDialog(store,cash_key,car,cashbook);capture('-vehicle-payments',obligations);obligations.reject()
                            from .household_categories_ui import CategoryRulesDialog
                            rules=CategoryRulesDialog(store,cash_key,cashbook);capture('-classification',rules);rules.reject()
                            from .home_ui import ScheduleDialog
                            schedule=ScheduleDialog(store,today,parent=window);schedule.title.setText('가족 생일 (예시)');schedule.kind.setCurrentIndex(schedule.kind.findData('birthday'));schedule.repeat.setCurrentIndex(schedule.repeat.findData('yearly'));schedule.alarm.setChecked(True)
                            capture('-schedule',schedule);schedule.close()
                            from .alert_ui import ScheduleAlertDialog
                            alert=ScheduleAlertDialog([{'date':today,'time':'09:00','title':'결혼기념일 · 예시','notes':'함께할 저녁 시간을 준비해 보세요.','counter':'결혼 10주년 · 함께한 날들을 기억해요'}],window);capture('-alerts',alert);alert.close()
                            from .email_ui import MailSettingsDialog
                            mail=MailSettingsDialog(store,window);capture('-mail-settings',mail);mail.close()
                            cashbook.close();assert not any(cash_key)
                            window.open_page("settings")
                            from .settings_defaults import SettingsDefaultsDialog
                            defaults=SettingsDefaultsDialog(window.current_page);capture('-settings-defaults',defaults);defaults.reject()
                            from .factory_reset_ui import BusinessResetDialog
                            reset_screen=BusinessResetDialog(window);capture('-business-reset',reset_screen);reset_screen.reject()
                            assert window.current_page.backup_grid.rowCount() >= 2
                            assert window.current_page.backup_exit.isChecked()
                            capture("-backup", window)
                            window.open_page("payments")
                            capture("-payments", window)
                            # Current release proof uses the actual new views,
                            # reauthentication dialog and manual-card ledger.
                            from .payment_access_ui import PaymentAccessDialog
                            from decimal import Decimal
                            payment_access=PaymentAccessDialog(window)
                            capture('-payment-access',payment_access)
                            payment_access.password.setText(DEMO_PASSWORD);payment_access.submit()
                            deadline=monotonic()+15
                            while payment_access.worker.isRunning() and monotonic()<deadline:
                                app.processEvents();sleep(.005)
                            app.processEvents();assert payment_access.verified
                            payment_access.deleteLater()
                            payment_supplier=store.trade_partners('supplier')[0]['id']
                            card_id=store.settle_supplier(payment_supplier,'15.25',date.today().isoformat(),'card','가상 수기 카드 기록 · 실제 승인 없음')
                            card_before=store.supplier_settlement_detail(card_id)
                            replacement=store.replace_supplier_payment('settlement',card_id,'10.25','card','가상 정정 결과','화면·원자성 검증',{key:card_before[key] for key in ('date','amount','method','memo')})
                            assert store.supplier_settlement_detail(card_id)['voided'] and store.supplier_settlement_detail(replacement)['amount']==Decimal('10.25')
                            window.current_page.refresh();window.current_page.tabs.setCurrentIndex(2)
                            capture('-payment-card-history',window)
                            # The frozen build must also render the document-
                            # scoped manager and persist a real split payment.
                            from .purchase_payment_ui import PurchasePaymentDialog
                            selected_purchase=next(r for r in store.documents(kind='purchase',status='posted')
                                if store.purchase_payment_context(r['id'])['cap']>=1)
                            direct_payment=PurchasePaymentDialog(window,selected_purchase['id'])
                            capture('-purchase-payment',direct_payment)
                            direct_state=direct_payment.snapshot
                            direct_result=store.add_purchase_payments(selected_purchase['id'],
                                {'cash':'0.25','transfer':'0.25','card':'0.50'},date.today().isoformat(),
                                '가상 분할 지급 · 송금/승인 없음',direct_state['state'])
                            assert len(direct_result['payment_ids'])==3
                            assert store.document(selected_purchase['id'])['cash_paid']==selected_purchase['cash_paid']+1
                            direct_payment.changed=True;direct_payment.refresh(True)
                            direct_payment.tabs.setCurrentIndex(1)
                            capture('-purchase-payment-history',direct_payment)
                            from .navigation import mark_saved
                            mark_saved(direct_payment);direct_payment.reject()
                            print('PASS direct purchase payment: real document manager, cash/transfer/manual-card split, ledger update and preserved history',flush=True)
                            # Exercise the newly packaged draft-plan and batch
                            # correction code, not just the legacy payment API.
                            from .purchase_payment_plan import normalize_plan
                            plan=normalize_plan(dict(partner_id=payment_supplier,target='10.25',
                                components=dict(cash=1,transfer=1,card=3,bill=1,dc=1,rebate=1,coupon='0.50',other_discount='1.75'),
                                datetime=date.today().isoformat()+' 16:12:00',account='ISOLATED-SMOKE-ACCOUNT',memo='격리 상세결제 검사',
                                bill=dict(number='ISOLATED-PLAN-BILL',due_date=date.today().isoformat(),bank='가상 은행')))
                            product=store.catalog_product(store.document(selected_purchase['id'])['lines'][0]['product_id'])
                            payload=dict(kind='purchase',date=date.today().isoformat(),partner_id=payment_supplier,paid=0,
                                lines=[dict(product_id=product['id'],qty=1,unit_price=str(Decimal('10.25')*product.get('price_basis',1)))],payment_plan=plan)
                            draft=store.save_draft(payload)
                            assert not store.document(draft)['payments']
                            store.post_document(payload,draft)
                            planned=store.document(draft);assert planned['balance']==0 and len(planned['payments'])==8
                            batch=store.purchase_payment_batch(planned['payments'][0]['id'])
                            assert batch['plan']==plan and batch['editable']
                            detailed=PurchasePaymentDialog(window,draft,editing=batch)
                            capture('-purchase-detailed-payment',detailed);mark_saved(detailed);detailed.reject()
                            correction=normalize_plan(dict(plan,target='6.00',components=dict(cash=5,dc=1),bill=None))
                            result=store.replace_purchase_payment_batch(batch['payment_ids'][0],correction,batch['state'],'격리 수정 검사')
                            assert store.document(draft)['balance']==Decimal('4.25')
                            latest=store.purchase_payment_batch(result['records'][0]['id'])
                            store.void_purchase_payment_batch(latest['payment_ids'][0],latest['state'],'격리 취소 검사')
                            assert store.document(draft)['balance']==Decimal('10.25')
                            print('PASS detailed purchase payment plan: draft without ledger writes, atomic eight-method posting, datetime/account, batch correction and cancellation',flush=True)
                            # The frozen executable exercises every finance
                            # view, plus a real non-cash batch and bill payout.
                            window.open_page('payments');finance_page=window.current_page
                            finance_state=store.payment_context(payment_supplier,date.today().isoformat())
                            finance_result=store.register_supplier_payment(payment_supplier,
                                {'transfer':'0.25','bill':'0.25','dc':'0.25'},date.today().isoformat(),
                                '격리 검사 · 실제 거래 아님',finance_state['state'],bill={
                                    'number':'ISOLATED-SMOKE-BILL','due_date':date.today().isoformat(),'bank':'가상 은행'})
                            assert finance_result['discount_total']==Decimal('0.25')
                            note=next(r for r in store.payable_bills(date.today().isoformat(),date.today().isoformat()) if r['number']=='ISOLATED-SMOKE-BILL')
                            debt=store.supplier_outstanding(payment_supplier)['balance']
                            store.pay_bill(note['id'],date.today().isoformat(),'transfer','격리 만기 지급')
                            assert store.supplier_outstanding(payment_supplier)['balance']==debt
                            store.add_business_entry(date.today().isoformat(),'expense','가상 업무 비용','0.25','transfer','검사 전용')
                            for index,suffix in ((0,'payables'),(1,'supplier-ledger'),(2,'payment-history'),(4,'bills'),(5,'receivables'),(6,'member-ledger'),(7,'sales-payments'),(8,'business-journal'),(9,'profit')):
                                finance_page.tabs.setCurrentIndex(index);finance_page.refresh();capture('-finance-'+suffix,window)
                            print('PASS finance workspace: 6 numbered menus, grouped 3D artwork, 9 real views, discount/bill batch and maturity payout without duplicate debt deduction',flush=True)
                            for page_key in ('purchase_suppliers','purchase_categories','purchase_products','purchase_documents','purchase_statement'):
                                assert window.open_page(page_key)
                                capture('-'+page_key.replace('_','-'),window)
                            window.open_page('members');capture('-customer-query',window)
                            from .appearance import apply_theme
                            for skin in ('spring','summer','autumn','winter'):
                                apply_theme(skin);capture('-skin-'+skin,window)
                            apply_theme('default')
                            print('PASS reference workspaces: customer filters, five purchase reports, payment authentication, atomic manual-card correction and four skins',flush=True)
                            window.open_page("permissions")
                            sample_staff = store.save_user({"code": "STAFF-SMOKE", "name": "권한 예시 직원", "role": "staff", "password": "Smoke-Staff-2026!", "permissions": ["catalog", "trade", "deny:products.delete", "deny:purchases.delete", "deny:payments.delete"]})
                            window.current_page.refresh()
                            window.current_page.user.setCurrentIndex(window.current_page.user.findData(sample_staff))
                            capture("-permissions", window)
                            fresh_login_store = Store(directory / "smoke-login.db")
                            try:
                                first_login = LoginDialog(fresh_login_store, window)
                                assert first_login.code.text() == "admin" and "1234" in first_login.hint.text()
                                capture("-login", first_login)
                                first_login.new_password.setText('Smoke-Admin-2026!');first_login.new_confirm.setText('Smoke-Admin-2026!')
                                first_login.submit()
                                assert fresh_login_store.actor["code"] == "admin"
                                assert not fresh_login_store.default_login_available()
                                first_login.close()
                                from .lucky_audio import login_sound_settings
                                from PySide6.QtWidgets import QGroupBox,QScrollArea
                                audio_capture_errors=[]
                                def capture_audio_settings():
                                    modal=app.activeModalWidget()
                                    try:
                                        assert modal is not None and len(modal.findChildren(QGroupBox))==6
                                        capture('-audio-settings',modal)
                                        modal.findChild(QScrollArea).verticalScrollBar().setValue(modal.findChild(QScrollArea).verticalScrollBar().maximum())
                                        capture('-audio-settings-more',modal)
                                    except BaseException as exc:audio_capture_errors.append(exc)
                                    finally:
                                        if modal is not None:modal.accept()
                                QTimer.singleShot(50,capture_audio_settings)
                                login_sound_settings(window)
                                assert not audio_capture_errors,audio_capture_errors
                            finally:
                                fresh_login_store.close()
                            from uuid import uuid4
                            stock_example=store.save_product({'name':'재고 흐름 검증용 상품 (예시)','base_unit':'g','price_basis':100,'purchase_price':'1650.00','sale_price':'2500.00','tracked':True})
                            store.post_stock_adjustment({'kind':'opening','product_id':stock_example,'date':today,'qty':'10000','unit_price':'1650.00','reason':'가상 기초재고'},uuid4().hex)
                            store.post_stock_adjustment({'kind':'waste','product_id':stock_example,'date':today,'qty':'500','unit_price':'0','reason':'화면 검증용 가상 폐기'},uuid4().hex)
                            for key,suffix in [('opening_stock','-opening-stock'),('stock_waste','-stock-waste'),('inventory_assets','-inventory-assets')]:
                                window.open_page(key);capture(suffix,window)
                            from .loyalty_ui import MemberDialog
                            member_detail=MemberDialog(store,parent=window);capture('-member-details',member_detail);member_detail.close()
                            from .email_ui import EmailCodeDialog
                            mail_recovery=EmailCodeDialog(store,'bind',window);capture('-email-recovery',mail_recovery);mail_recovery.close()
                            window.popup_menus=True
                            window.open_page('products');window.open_page('purchases');window.open_page('worklog');app.processEvents()
                            window.pages['worklog'].tabs.setCurrentIndex(0);app.processEvents()
                            assert all(not popup.isModal() for popup in window.popup_windows.values())
                            assert all(window.scrollers[key].isVisible() and window.pages[key].isVisible() for key in window.popup_windows)
                            assert window.pages['worklog'].grid.isVisible()
                            capture('-worklog-independent',window.popup_windows['worklog'])
                            capture('-independent-windows',window.popup_windows['purchases'])
                            popup=window.popup_windows['worklog'];popup.resize(720,520);app.processEvents()
                            bar=popup.query_bar;position=bar.pos()
                            window.scrollers['worklog'].verticalScrollBar().setValue(window.scrollers['worklog'].verticalScrollBar().maximum())
                            app.processEvents();assert bar.isVisible() and bar.pos()==position and popup.isSizeGripEnabled()
                            capture('-resized-fixed-query',popup)
                            popup.resize(1180,800)
                            from .private_session import remember_key
                            single_key=store.unlock_diary('Smoke-Diary-Only-2026!');remember_key(store,'diary',single_key);single_key[:]=bytes(len(single_key))
                            window.open_page('diary');app.processEvents()
                            assert 'diary' not in window.popup_windows
                            assert window.personal_windows['diary'].isVisible()
                            capture('-diary-independent',window.personal_windows['diary'])
                            from .independent_windows import close_independent_windows
                            assert close_independent_windows(window)
                            window.popup_menus=False;window.open_page('home')
                            from .lucky_numbers import open_lucky_numbers
                            lucky=open_lucky_numbers(window)
                            lucky.page.draw_five()
                            for _ in range(30):lucky.page.finish_draw()
                            assert len(lucky.page.results)==5 and all(len(set(values))==6 for values in lucky.page.results)
                            capture('-lucky-numbers',lucky);lucky.close()
                            lucky.page.reset(False);lucky.show();lucky.page.draw_six()
                            for _ in range(120):lucky.page.machine.motion.advance(1/60,True)
                            capture('-lucky-mixing',lucky);lucky.close()
                            # New paper workflows use the same saved snapshot in
                            # the editor, native preview, PNG and email payload.
                            window.open_page('store_profile');capture('-store-profile',window)
                            from .offline_ui import SalesFormDialog
                            customer=store.trade_partners('customer')[0]
                            sample_product=store.trade_products()[0]
                            order_id=store.save_sales_form('order',dict(date=today,partner_id=customer['id'],memo='인수 시 수량·규격 확인',lines=[dict(product_id=sample_product['id'],qty='2',unit_price=str(sample_product['sale_price']))],order=dict(delivery_date=today,assignee='담당자',phone='010-0000-0000',address='점포 앞 · 예시',delivery_note='도착 전 연락',order_type='phone')))
                            window.open_page('orders');capture('-orders',window)
                            order_editor=SalesFormDialog(window,'orders',store.sales_form(order_id),parent=window);capture('-order-editor',order_editor);order_editor.close()
                            from .commercial_paper import paper_document,paper_images
                            paper=paper_document(store.sales_form(order_id),'주문서')
                            assert next(paper_images(paper)).save(str(args.screenshot.with_name(args.screenshot.stem+'-paper-a4.png')))
                            from PySide6.QtPrintSupport import QPrinter,QPrintPreviewDialog
                            printer=QPrinter();printer.setOutputFormat(QPrinter.OutputFormat.PdfFormat)
                            preview=QPrintPreviewDialog(printer,window);preview.paintRequested.connect(paper.print_);capture('-commercial-preview',preview);preview.close()
                            assert not any('오락실' in button.text() for button in lucky.page.findChildren(QPushButton))
                            # Synthetic catalogue only. Never package uploaded
                            # supplier/price lists or personal business records.
                            store.seed_catalog_categories();window.open_page('products');capture('-catalog-tree-tax',window.popup_windows.get('products',window))
                            # Both supplier controls use the same catalogue
                            # filter. Exercise the real frozen Qt popup, using
                            # synthetic supplier metadata only.
                            catalog_page=window.pages['products'];catalog_page.refresh()
                            catalog_page.supplier.showPopup();app.processEvents()
                            capture('-catalog-supplier-dropdown',catalog_page.supplier.view().window())
                            catalog_page.supplier.hidePopup()
                            from .catalog_ui import SupplierFilterDialog
                            supplier_filter=SupplierFilterDialog(store,window,catalog_page._catalog_suppliers)
                            filter_supplier=catalog_page._catalog_suppliers[0]
                            supplier_filter.search.setText(filter_supplier['code']);supplier_filter.refresh()
                            assert supplier_filter.grid.selected()['id']==filter_supplier['id']
                            capture('-catalog-supplier-search',supplier_filter)
                            supplier_filter.choose()
                            assert supplier_filter.result()==QDialog.DialogCode.Accepted and supplier_filter.record['id']==filter_supplier['id']
                            catalog_page.supplier.setCurrentIndex(catalog_page.supplier.findData(filter_supplier['id']))
                            assert all(r['supplier_id']==filter_supplier['id'] for r in catalog_page._catalog_matches)
                            catalog_page.supplier.setCurrentIndex(0)
                            print('PASS catalogue supplier: real dropdown, name/code popup, selection and reset',flush=True)
                            from decimal import Decimal
                            record=store.products()[0];store.save_product(dict(sale_price=str(record['sale_price']+Decimal('0.25'))),record['id'])
                            member_id=store.save_member(dict(code='SMOKE-FIXED',name='지정가 검증 회원'));store.save_member_price(dict(member_id=member_id,product_id=record['id'],price='123.45'))
                            editor=ProductDialog(store,next(r for r in store.products() if r['id']==record['id']),window);capture('-product-tax-details',editor)
                            editor.tabs.setCurrentIndex(4);capture('-product-change-log',editor);editor.tabs.setCurrentIndex(5);capture('-product-member-prices',editor);editor.close()
                            from openpyxl import Workbook
                            import xlrd
                            assert xlrd.__version__=='2.0.2'
                            workbook=Workbook();sheet=workbook.active;sheet.append(['상품명','규격','바코드','공급사','매입가','판매가'])
                            sheet.append(['생삼겹살','g','SAMPLE-RAW','예시 공급사','1000.25','2000.50']);sheet.append(['양념갈비','팩','SAMPLE-COOKED','예시 공급사','2000.25','3000.50']);sheet.append(['분류 검토 상품','','SAMPLE-REVIEW','예시 공급사','0','0'])
                            excel=directory/'synthetic-catalog.xlsx';workbook.save(excel)
                            from .catalog_import import plan_excel
                            from .catalog_import_ui import CatalogImportDialog
                            importer=CatalogImportDialog(store,window);importer.loaded(plan_excel(excel));capture('-catalog-excel-preview',importer);importer.close()
                            # Windows smoke uses real DPAPI, not a test cipher.
                            from .unsaved_work import CrashJournal,RecoveryCenter
                            window.navigation.enable();journal=CrashJournal(window);window.crash_journal=journal
                            window.open_page('products');unsaved=ProductDialog(store,parent=window);unsaved.show();app.processEvents();unsaved.fields['name'].setText('임시 복구 검증 상품');journal.capture(unsaved)
                            if sys.platform=='win32':
                                rows=journal.rows();assert len(rows)==1
                                assert journal.read(rows[0]['identifier'])['values']
                                assert '임시 복구 검증 상품'.encode() not in next(journal.root.glob('*.draft')).read_bytes()
                            recovery=RecoveryCenter(journal);capture('-crash-recovery',recovery);recovery.accept()
                            journal.forget(unsaved);window.navigation.suspended=True;unsaved.close();journal.stop();window.navigation.disable();window.navigation.suspended=False
                            # Exercise the real close event and its real modal
                            # dialog in the frozen EXE, not a mocked question.
                            window._backup_runtime_enabled=True
                            def cancel_exit():
                                box=QApplication.activeModalWidget()
                                assert isinstance(box,QMessageBox) and '종료하시겠습니까?' in box.text()
                                assert box.defaultButton()==box.button(QMessageBox.StandardButton.No)
                                capture('-exit-confirmation',box);box.close()
                            QTimer.singleShot(0,cancel_exit)
                            assert not window.close() and window.isVisible() and not window._exit_confirmed
                            window._backup_runtime_enabled=False
                            # Only this isolated temporary smoke DB is reset.
                            # This password is synthetic, never an application
                            # default or an uploaded user's reset credential.
                            from .factory_reset import private_digest
                            smoke_reset_password='Synthetic-Reset-Smoke-Only!'
                            store.configure_reset_authority(DEMO_PASSWORD,'wonku9',smoke_reset_password,smoke_reset_password)
                            preserved=private_digest(store.db)
                            reset_result=store.reset_business_data(DEMO_PASSWORD,'wonku9',smoke_reset_password,'업무자료 초기화')
                            assert reset_result['private_preserved'] and Path(reset_result['safety_backup']).is_file()
                            assert private_digest(store.db)==preserved
                            store.login(DEMO_CODE,DEMO_PASSWORD)
                            assert all(count==0 for count in store.business_reset_plan().values())
                            print('PASS business reset: protected backups, exact private preservation, empty business data and existing login',flush=True)
                    if args.smoke_test:
                        print(f"PASS standalone-v2: {len(window.items)} desktop pages, returns, expanded units, persistence, backup restore, no legacy imports", flush=True)
                        app.exit(0)
                except BaseException:
                    import traceback
                    traceback.print_exc()
                    app.exit(1)
            QTimer.singleShot(250, verify)
        return app.exec()
    except BaseException as exc:
        if args.smoke_test:
            import traceback
            traceback.print_exc()
        else:
            QMessageBox.critical(None, "독립판 실행 오류", f"{exc}\n\n새 자료 위치: {path}\n기존 프로그램의 데이터 파일을 지정하지 마세요.")
        return 1
    finally:
        if window:
            window.close()
            window.deleteLater()
        if backdrop:
            backdrop.close()
        if store:
            store.close()
        if process_lock:
            process_lock.unlock()
        if scratch:
            scratch.cleanup()
