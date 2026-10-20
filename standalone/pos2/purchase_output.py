"""Purchase-voucher forms: one selected voucher, never the whole query list.

Prices, names, quantities and tax modes come from saved document lines.
Legacy vouchers do not snapshot contact details/codes/specifications, so those
supplemental current registration fields are explicitly labelled as such.
"""
from collections import defaultdict
from decimal import Decimal
from html import escape
from io import BytesIO
from pathlib import Path

from PySide6.QtCore import QIODevice,QSaveFile,QSizeF
from PySide6.QtGui import QPageSize,QTextDocument
from PySide6.QtPrintSupport import QPrinter,QPrintPreviewDialog
from PySide6.QtWidgets import QFileDialog,QMessageBox

from .font_support import regular_font
from .widgets import DataGrid,authorize_output

HEADERS=('No.','코드 / 바코드','품명','규격','단위','수량','단가','금액','과면세')
NOTICE='업체·코드·규격은 현재 등록정보입니다. 저장 전표의 품명·단위·수량·단가·금액·과면세는 거래 당시 값을 사용합니다.'
STATUSES={'posted':'확정','draft':'초안','cancelled':'취소'}


def purchase_record(dialog):
    store=dialog.store;store.require_menu('purchases')
    products={r['id']:r for r in store.products()}
    if dialog.readonly:
        record=store.document(dialog.record['id'])
        if not record['kind'].startswith('purchase'):raise ValueError('매입전표를 선택하세요.')
        lines=[{**line,'qty':Decimal(line['input_qty'])} for line in record['lines']]
    else:
        dialog.recalculate();lines=[]
        from .common import decimal as number
        for row in dialog.line_rows:
            product=products.get(row['product'].currentData())
            if not product:continue
            quantity=number(row['qty'].text())
            if dialog.is_return and quantity==0:continue
            price=number(row['price'].text().replace(',',''))
            total=number(row['total'].text().replace(',',''))
            if quantity<=0:raise ValueError('출력할 수량은 0보다 커야 합니다.')
            lines.append(dict(product_id=product['id'],product_name=product['name'],unit_name=row['unit'].currentText(),qty=quantity,unit_price=price,total=total,tax_mode=product['tax_mode'],price_basis=row.get('price_basis',1)))
        record=dict(kind=dialog.kind,number=(dialog.record or {}).get('number','미확정 입력'),date=dialog.date.date().toString('yyyy-MM-dd'),partner_id=dialog.partner.currentData(),status='draft',memo=dialog.memo.text(),user_name=store.actor['name'],total=sum((Decimal(str(l['total'])) for l in lines),Decimal(0)),cash_paid=Decimal(0),balance=Decimal(0))
    if not lines:raise ValueError('출력할 상품이 없습니다.')
    supplier=next((p for p in store.partners('supplier') if p['id']==record['partner_id']),None)
    if supplier is None:raise ValueError('전표의 공급사를 확인하세요.')
    for line in lines:
        current=products.get(line['product_id'],{})
        line.update(code=current.get('barcode') or current.get('code',''),specification=current.get('specification',''))
    title='매 입 반 품 전 표' if record['kind']=='purchase_return' else '매 입 전 표'
    return {**record,'title':title,'supplier':supplier,'recipient':store.store_profile(),'lines':lines,'unsaved':not dialog.readonly}


def quantity_text(record):
    totals=defaultdict(Decimal)
    for line in record['lines']:totals[line['unit_name']]+=Decimal(str(line['qty']))
    return ' / '.join(f'{value:,.3f}'.rstrip('0').rstrip('.')+' '+unit for unit,value in totals.items())


def line_values(index,line):
    unit=line['unit_name']+(' (100g당)' if line.get('price_basis')==100 else '')
    return [index,line['code'],line['product_name'],line['specification'],unit,Decimal(str(line['qty'])),Decimal(str(line['unit_price'])),Decimal(str(line['total'])),'과' if line['tax_mode']=='taxable' else '면']


def parties(record):
    supplier,recipient=record['supplier'],record['recipient']
    return [(label,supplier.get(a,''),recipient.get(b,'')) for label,a,b in (
        ('사업자번호','business_no','business_no'),('상호','name','store_name'),
        ('대표자','representative','representative'),('주소','address','address'),
        ('업태','business_type','business_type'),('종목','business_item','business_category'),('전화','phone','phone'))]


def status_text(record):
    status=STATUSES[record['status']]
    return status+(' · 미저장 입력 미리보기 / 재고·결제 미반영' if record.get('unsaved') else ' · 유효 거래로 사용하지 마세요' if record['status']=='cancelled' else '')


def purchase_html(record):
    e=lambda value:escape(str(value if value is not None else ''))
    pages=[];lines=record['lines']
    # Explicit short chunks repeat both parties and keep A4 forms readable.
    for offset in range(0,len(lines),18):
        chunk=lines[offset:offset+18]
        rows=[]
        for index,line in enumerate(chunk,offset+1):
            values=line_values(index,line)
            def display(column,value):
                if not isinstance(value,Decimal):return value
                return f'{value:,.3f}'.rstrip('0').rstrip('.') if column==5 else f'{value:,.2f}'
            rows.append('<tr>'+''.join('<td>'+e(display(column,value))+'</td>' for column,value in enumerate(values))+'</tr>')
        rows+=['<tr><td colspan="9">&nbsp;</td></tr>']*(18-len(chunk))
        party_rows=''.join(f'<tr><th width="12%">{label}</th><td width="38%">{e(left)}</td><th width="12%">{label}</th><td width="38%">{e(right)}</td></tr>' for label,left,right in parties(record))
        heading=' style="page-break-before:always"' if offset else ''
        pages.append(f'''<h1 align="center"{heading}>{e(record['title'])}</h1>
        <p>매입일: {e(record['date'])}　전표번호: {e(record['number'])}　{e(status_text(record))}</p>
        <table width="100%"><tr><th colspan="2">공 급 자</th><th colspan="2">공 급 받 는 자</th></tr>{party_rows}</table><br>
        <table width="100%"><thead><tr>{''.join('<th>'+e(v)+'</th>' for v in HEADERS)}</tr></thead>{''.join(rows)}</table>
        <p align="right">페이지 소계: {sum(Decimal(str(l['total'])) for l in chunk):,.2f} 원　 /　전표 총합계: {Decimal(str(record['total'])):,.2f} 원</p>
        <p>단위별 총수량: {e(quantity_text(record))}<br>메모: {e(record.get('memo',''))}<br>
        담당자: {e(record.get('user_name',''))}　품목 묶음 {offset//18+1}/{(len(lines)+17)//18}<br>{e(NOTICE)}<br>과/면은 저장된 구분이며 이 문서는 법정 세금계산서가 아닙니다.</p>''')
    return '<html><head><style>body{font-size:8pt;color:#111}h1{font-size:22pt}table{border-collapse:collapse}td,th{border:1px solid #333;padding:3px}th{background:#E3E3E3}</style></head><body>'+''.join(pages)+'</body></html>'


def purchase_document(record):
    document=QTextDocument();document.setDefaultFont(regular_font(8));document.setDocumentMargin(28)
    document.setPageSize(QSizeF(794,1123));document.setHtml(purchase_html(record));return document


def purchase_workbook(record):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment,Border,Font,PatternFill,Side
    from openpyxl.utils import get_column_letter
    book=Workbook();sheet=book.active;sheet.title='매입전표';sheet.sheet_view.showGridLines=False
    border=Border(**dict.fromkeys(('left','right','top','bottom'),Side(style='thin',color='333333')))
    def put(row,start,end,value,bold=False,fill=False,align='left'):
        if end>start:sheet.merge_cells(start_row=row,start_column=start,end_row=row,end_column=end)
        cell=sheet.cell(row,start,value)
        if isinstance(value,str):cell.data_type='s'  # Never execute user spreadsheet formulas.
        cell.font=Font(name='맑은 고딕',size=10,bold=bold)
        cell.alignment=Alignment(horizontal=align,vertical='center',wrap_text=True)
        for column in range(start,end+1):
            item=sheet.cell(row,column);item.border=border
            if fill:item.fill=PatternFill('solid',fgColor='E3E3E3')
        return cell
    put(1,1,10,record['title'],True,align='center').font=Font(name='맑은 고딕',size=22,bold=True)
    put(2,1,10,f"매입일: {record['date']}    전표번호: {record['number']}    {status_text(record)}")
    put(3,1,5,'공 급 자',True,True,'center');put(3,6,10,'공 급 받 는 자',True,True,'center')
    for row,(label,left,right) in enumerate(parties(record),4):
        put(row,1,1,label,True);put(row,2,5,left);put(row,6,6,label,True);put(row,7,10,right)
        sheet.row_dimensions[row].height=max(23,15*(1+max(len(str(left)),len(str(right)))//38))
    columns=((1,1),(2,2),(3,4),(5,5),(6,6),(7,7),(8,8),(9,9),(10,10))
    for (start,end),text in zip(columns,HEADERS):put(11,start,end,text,True,True,'center')
    for index in range(max(18,len(record['lines']))):
        row=12+index;sheet.row_dimensions[row].height=28
        values=line_values(index+1,record['lines'][index]) if index<len(record['lines']) else ['']*9
        for (start,end),value in zip(columns,values):
            numeric=isinstance(value,Decimal)
            cell=put(row,start,end,float(value) if numeric else value,align='right' if numeric else 'left')
            if numeric:cell.number_format='#,##0.###' if start==7 else '#,##0.00'
    bottom=12+max(18,len(record['lines']))
    put(bottom,1,8,'전표 총합계',True,True,'right');put(bottom,9,10,float(record['total']),True).number_format='#,##0.00'
    for row,value in enumerate((f"단위별 총수량: {quantity_text(record)}",f"메모: {record.get('memo','')}",f"담당자: {record.get('user_name','')}",NOTICE,'법정 세금계산서가 아닙니다.'),bottom+1):
        put(row,1,10,value);sheet.row_dimensions[row].height=30
    for i,width in enumerate((12,20,15,10,14,14,11,15,17,8),1):sheet.column_dimensions[get_column_letter(i)].width=width
    sheet.row_dimensions[1].height=36;sheet.row_dimensions[2].height=30
    sheet.freeze_panes='C12';sheet.print_title_rows='1:11';sheet.print_area=f'A1:J{bottom+5}'
    sheet.sheet_properties.pageSetUpPr.fitToPage=True
    sheet.page_setup.orientation='portrait';sheet.page_setup.paperSize=sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth=1;sheet.page_setup.fitToHeight=0
    sheet.print_options.horizontalCentered=True
    sheet.oddFooter.center.text='&P / &N';sheet.oddFooter.right.text='SY POS'
    return book


def atomic_write(path,data):
    output=QSaveFile(str(path))
    if not output.open(QIODevice.OpenModeFlag.WriteOnly) or output.write(data)!=len(data) or not output.commit():raise OSError('파일을 저장하지 못했습니다. 저장 위치와 디스크 공간을 확인하세요.')


def output_purchase(dialog,mode):
    if not authorize_output(dialog):return
    record=purchase_record(dialog)
    if mode in {'xlsx','pdf'}:
        suffix='.'+mode
        name,_=QFileDialog.getSaveFileName(dialog,'매입전표 '+mode.upper()+' 저장',record['number']+suffix,'Excel 통합문서 (*.xlsx)' if mode=='xlsx' else 'PDF 문서 (*.pdf)')
        if not name:return
        if not authorize_output(dialog):return
        record=purchase_record(dialog)  # Refresh after the native chooser's event loop.
        path=Path(name).with_suffix(suffix)
        if path!=Path(name) and path.exists() and QMessageBox.question(dialog,'파일 덮어쓰기',f'{path.name} 파일을 교체할까요?',QMessageBox.StandardButton.Yes|QMessageBox.StandardButton.No,QMessageBox.StandardButton.No)!=QMessageBox.StandardButton.Yes:return
        if mode=='xlsx':
            buffer=BytesIO();purchase_workbook(record).save(buffer)
            if not authorize_output(dialog):return
            atomic_write(path,buffer.getvalue())
        else:
            # Render to a temporary file first; a failed output must not truncate
            # the user's previous PDF. QPrinter needs a filesystem destination.
            from tempfile import TemporaryDirectory
            with TemporaryDirectory(prefix='sy-purchase-') as tmp:
                target=Path(tmp)/'voucher.pdf';printer=QPrinter(QPrinter.PrinterMode.HighResolution)
                printer.setPageSize(QPageSize(QPageSize.PageSizeId.A4));printer.setOutputFormat(QPrinter.OutputFormat.PdfFormat);printer.setOutputFileName(str(target))
                document=purchase_document(record)
                if not authorize_output(dialog):return
                document.print_(printer)
                data=target.read_bytes()
                if not data.startswith(b'%PDF'):raise ValueError('PDF를 생성하지 못했습니다.')
                if not authorize_output(dialog):return
                atomic_write(path,data)
        return path
    if mode not in {'grid','preview'}:raise ValueError('출력 방식을 확인하세요.')
    printer=QPrinter(QPrinter.PrinterMode.HighResolution);printer.setPageSize(QPageSize(QPageSize.PageSizeId.A4))
    preview=QPrintPreviewDialog(printer,dialog);preview.setWindowTitle(record['title']+' · '+('인쇄그리드' if mode=='grid' else '프린터 양식')+' 미리보기');preview.resize(1100,800)
    def paint(target):
        if not authorize_output(dialog):preview.reject();return
        try:
            fresh=purchase_record(dialog)
            if mode=='grid':
                grid=DataGrid(HEADERS);grid.setSortingEnabled(False)
                grid.fill([line_values(i,line) for i,line in enumerate(fresh['lines'],1)])
                document=QTextDocument();document.setDefaultFont(regular_font(8))
                document.setHtml(grid.document_html(fresh['title'],f"{fresh['number']} · {fresh['date']} · {status_text(fresh)}\n공급자: {fresh['supplier']['name']} / 공급받는 자: {fresh['recipient']['store_name']}\n합계 {fresh['total']:,.2f} 원"));grid.deleteLater()
            else:document=purchase_document(fresh)
            if not authorize_output(dialog):preview.reject();return
            document.print_(target)
        except (ValueError,PermissionError,ConnectionError) as exc:
            preview.reject();QMessageBox.warning(dialog,'매입전표 출력',str(exc))
    preview.paintRequested.connect(paint);preview.exec()
