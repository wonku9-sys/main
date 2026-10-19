"""Exercise manual-inspired pages with the actual local SQLite implementation."""
import os
from datetime import date
from decimal import Decimal
from types import SimpleNamespace

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import pytest
from openpyxl import load_workbook
from PySide6.QtCore import QCoreApplication, QDate, QEvent, Qt
from PySide6.QtWidgets import QApplication, QFileDialog, QMessageBox

from seoyoung_pos.app import STYLE
from seoyoung_pos.offline_pages import (
    AnalysisPage, LedgerPage, MemberDialog, MembersPage, PointsDialog,
    PromotionDialog, PromotionsPage, StocktakePage,
)
from seoyoung_pos.service import POSService


@pytest.fixture(scope="module")
def application():
    result = QApplication.instance() or QApplication([])
    result.setStyle("Fusion")
    result.setStyleSheet(STYLE)
    yield result
    for widget in result.topLevelWidgets():
        widget.close()
        widget.deleteLater()
    QCoreApplication.sendPostedEvents(None, QEvent.Type.DeferredDelete)
    result.processEvents()


@pytest.fixture
def workspace(application, tmp_path, monkeypatch):
    messages = []
    monkeypatch.setattr(QMessageBox, "warning", lambda _parent, _title, text, *_args: messages.append(text))
    monkeypatch.setattr(QMessageBox, "information", lambda _parent, _title, text, *_args: messages.append(text))
    monkeypatch.setattr(QMessageBox, "question", lambda *_args, **_kwargs: QMessageBox.StandardButton.Yes)
    service = POSService(tmp_path / "offline.db")
    service.create_initial_admin("admin", "관리자", "Safe-pass-123")
    service.login("admin", "Safe-pass-123")
    supplier = service.save_partner({"kind": "supplier", "code": "S001", "name": "서영 공급사"})
    customer = service.save_partner({"kind": "customer", "code": "C001", "name": "매장 판매처"})
    product = service.save_product({"code": "P001", "name": "한돈 삼겹살", "unit": "kg", "category": "돈육", "supplier_id": supplier, "purchase_price": 10000, "sale_price": 15000})
    second = service.save_product({"code": "P002", "name": "포장팩", "unit": "EA", "category": "부자재", "supplier_id": supplier, "purchase_price": 200, "sale_price": 300})
    today = date.today().isoformat()
    service.create_document("purchase", supplier, today, [{"product_id": product, "qty": "10", "unit_price": 10000}, {"product_id": second, "qty": "20", "unit_price": 200}], paid=20000)
    service.create_document("sale", customer, today, [{"product_id": product, "qty": "2.5", "unit_price": 15000}], paid=10000)
    result = SimpleNamespace(service=service, show_message=messages.append, messages=messages,
                             product=product, second=second, supplier=supplier, customer=customer)
    yield result
    for widget in application.topLevelWidgets():
        widget.close()
        widget.deleteLater()
    QCoreApplication.sendPostedEvents(None, QEvent.Type.DeferredDelete)
    application.processEvents()
    service.close()


def test_ledger_real_charges_payments_and_date_validation(workspace):
    page = LedgerPage(workspace)
    page.refresh()
    assert page.table.rowCount() == 2
    assert "84,000" in page.total_label.text()
    page.kind.setCurrentIndex(page.kind.findData("customer"))
    assert "27,500" in page.total_label.text()
    assert page.table.rowCount() == 2
    page.start.setDate(QDate.currentDate().addDays(1))
    page.refresh()
    assert "시작일" in workspace.messages[-1]
    assert page.table.rowCount() == 2


def test_analysis_units_are_separate_and_excel_is_typed(workspace, monkeypatch, tmp_path):
    page = AnalysisPage(workspace)
    page.refresh()
    assert page.table.rowCount() == 2
    assert {page.table.item(row, 1).text() for row in range(2)} == {"kg", "EA"}
    assert "37,500" in page.total_label.text()
    assert "12,500" in page.total_label.text()
    for _title, mode in page.GROUPS:
        page.group.setCurrentIndex(page.group.findData(mode))
        assert page.table.rowCount() == 2
    target = tmp_path / "analysis.xlsx"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *_args, **_kwargs: (str(target), ""))
    page.export()
    sheet = load_workbook(target).active
    assert sheet.max_row == 3
    assert isinstance(sheet.cell(2, 5).value, int)
    assert isinstance(sheet.cell(2, 3).value, (int, float))


def test_stocktake_count_fields_survive_sort_and_stale_rejection(workspace):
    page = StocktakePage(workspace)
    page.refresh()
    page.count_fields[workspace.product].setText("8.125")
    page.count_fields[workspace.second].setText("0")
    page.table.sortItems(0, Qt.SortOrder.DescendingOrder)
    assert page.difference_labels[workspace.product].text() == "0.625"
    page.reason.setText("냉장창고 실사")
    page.save()
    products = {record["id"]: record for record in workspace.service.list_products()}
    assert products[workspace.product]["stock_qty"] == "8.125"
    assert products[workspace.second]["stock_qty"] == "0.000"
    assert page.history.rowCount() == 1
    assert all(not field.text() for field in page.count_fields.values())
    page.count_fields[workspace.product].setText("10")
    page.reason.setText("변경된 장부 실사")
    workspace.service.adjust_stock(workspace.product, "1", "실사 입력 중 입고")
    page.save()
    assert "장부 재고가 변경" in workspace.messages[-1]
    assert page.count_fields[workspace.product].text() == "10"
    assert len(workspace.service.list_stocktakes()) == 1


def test_stocktake_export_contains_widget_inputs(workspace, tmp_path, monkeypatch):
    page = StocktakePage(workspace)
    page.refresh()
    page.count_fields[workspace.product].setText("8")
    target = tmp_path / "stocktake.xlsx"
    monkeypatch.setattr(QFileDialog, "getSaveFileName", lambda *_args, **_kwargs: (str(target), ""))
    page.export()
    rows = list(load_workbook(target).active.values)
    row = next(row for row in rows if row[0] == "P001")
    assert row[3:6] == (7.5, 8, 0.5)


def test_member_registration_points_and_history(workspace):
    dialog = MemberDialog(workspace.service)
    dialog.fields["code"].setText("M001")
    dialog.fields["name"].setText("서영 고객")
    dialog.fields["phone"].setText("010-0000-0000")
    dialog.save()
    member = workspace.service.list_members()[0]
    adjustment = PointsDialog(workspace.service, member)
    adjustment.points.setText("1000")
    adjustment.reason.setText("수기 적립")
    adjustment.save()
    page = MembersPage(workspace)
    page.refresh()
    page.table.selectRow(0)
    assert page.history.rowCount() == 1
    assert page.history.item(0, 1).sort_value == 1000
    assert "1,000" in page.total_label.text()
    adjustment = PointsDialog(workspace.service, member)
    adjustment.points.setText("-300")
    adjustment.reason.setText("포인트 사용")
    adjustment.save()
    page.refresh()
    assert page.history.rowCount() == 2
    assert workspace.service.member_points(member["id"]) == 700
    page.delete()
    assert "포인트 이력" in workspace.messages[-1]
    assert len(workspace.service.list_members()) == 1


def test_promotion_form_real_price_and_overlap_validation(workspace):
    dialog = PromotionDialog(workspace.service)
    dialog.product.setCurrentIndex(dialog.product.findData(workspace.product))
    dialog.title.setText("주말 한돈 특가")
    dialog.sale_price.setText("12000")
    dialog.save()
    page = PromotionsPage(workspace)
    page.refresh()
    assert page.table.rowCount() == 1
    assert page.table.item(0, 7).sort_value == 12000
    assert page.table.item(0, 6).text() == "일반가"
    assert workspace.service.get_effective_price(workspace.product, date.today().isoformat()) == 12000
    duplicate = PromotionDialog(workspace.service)
    duplicate.product.setCurrentIndex(duplicate.product.findData(workspace.product))
    duplicate.title.setText("중복 특가")
    duplicate.sale_price.setText("11000")
    duplicate.save()
    assert "겹치는 특매" in workspace.messages[-1]
    assert len(workspace.service.list_promotions()) == 1


def test_dedicated_member_and_promotion_permissions(workspace):
    service = workspace.service
    service.save_user({"username": "loyalty", "name": "회원 담당", "role": "staff", "password": "Safe-pass-456", "permissions": ["members", "promotions"]})
    service.login("loyalty", "Safe-pass-456")
    members = MembersPage(workspace)
    members.refresh()
    assert members.editable
    promotions = PromotionsPage(workspace)
    promotions.refresh()
    assert promotions.editable
    dialog = PromotionDialog(service)
    assert dialog.product.count() == 2
    with pytest.raises(PermissionError):
        service.list_products()
