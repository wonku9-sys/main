"""Actual desktop navigation, period prices and business use with no network."""
import os
import socket
from datetime import date

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import pytest
from openpyxl import load_workbook
from PySide6.QtCore import QCoreApplication, QDate, QEvent
from PySide6.QtWidgets import QApplication, QMessageBox

from seoyoung_pos.app import MainWindow, NAV_ITEMS, STYLE
from seoyoung_pos.demo import populate_demo
from seoyoung_pos.dialogs import DocumentDialog
from seoyoung_pos.export import export_workbook
from seoyoung_pos.service import POSService


@pytest.fixture(scope="module")
def application():
    app = QApplication.instance() or QApplication([])
    app.setStyle("Fusion")
    app.setStyleSheet(STYLE)
    yield app
    for widget in app.topLevelWidgets():
        widget.close()
        widget.deleteLater()
    QCoreApplication.sendPostedEvents(None, QEvent.Type.DeferredDelete)
    app.processEvents()


@pytest.fixture
def workspace(application, tmp_path, monkeypatch):
    messages = []
    monkeypatch.setattr(QMessageBox, "warning", lambda _parent, _title, text, *_args: messages.append(text))
    monkeypatch.setattr(QMessageBox, "question", lambda *_args, **_kwargs: QMessageBox.StandardButton.Yes)
    service = POSService(tmp_path / "store.db")
    populate_demo(service)
    yield service, messages
    for widget in application.topLevelWidgets():
        widget.close()
        widget.deleteLater()
    QCoreApplication.sendPostedEvents(None, QEvent.Type.DeferredDelete)
    application.processEvents()
    service.close()


def test_menu_tree_toolbar_and_navigation_history_are_real(application, workspace):
    service, messages = workspace
    window = MainWindow(service, service.path, demo=True)
    window.show()
    assert len(window.nav_buttons) == len(NAV_ITEMS) == 15
    assert window.navigation_tree.topLevelItemCount() == 6
    assert not window.toolbar_actions["back"].isEnabled()
    for key in ("products", "purchases", "stocktake", "members", "analysis"):
        window.page_actions[key].trigger()
        application.processEvents()
        assert window.current_key == key
        assert window.navigation_tree.currentItem() is window.nav_items[key]
        assert window.nav_buttons[key].isChecked()
        assert window.page_actions[key].isChecked()
    window.toolbar_actions["back"].trigger()
    assert window.current_key == "members"
    window.toolbar_actions["forward"].trigger()
    assert window.current_key == "analysis"
    window.toolbar_actions["back"].trigger()
    window.navigate("ledger")
    assert not window.toolbar_actions["forward"].isEnabled()
    window.navigation_tree.setCurrentItem(window.nav_items["promotions"])
    assert window.current_key == "promotions"
    window.navigate("purchases")
    window.rebuild_navigation()
    assert window.current_key == "purchases"  # Page code is singular; navigation key isn't.
    assert len(window.navigation_history) == 1
    assert not window.toolbar_actions["back"].isEnabled()
    assert messages == []


def test_new_workflows_do_not_grant_existing_staff_extra_access(workspace):
    service, messages = workspace
    service.save_user({"username": "catalog", "name": "상품 담당", "password": "Staff-12345", "role": "staff", "permissions": ["products"]})
    service.login("catalog", "Staff-12345")
    window = MainWindow(service, service.path)
    assert set(window.nav_buttons) == set(window.page_actions) == {"products"}
    assert window.navigation_tree.topLevelItemCount() == 1
    assert window.navigate("members") is False
    assert window.navigate("promotions") is False
    assert window.current_key == "products"
    with pytest.raises(PermissionError):
        service.adjust_member_points(1, 100, "허용되지 않은 적립")
    assert messages == []


def test_period_price_uses_document_date_without_overwriting_manual_price(workspace):
    service, messages = workspace
    product = service.list_products()[0]
    partner = service.list_partners("customer")[0]
    base = product["sale_price"]
    service.save_promotion({"product_id": product["id"], "title": "기간행사", "date_from": "2026-10-01", "date_to": "2026-10-10", "sale_price": 12345})
    dialog = DocumentDialog(service, "sale", partner=partner)
    dialog.date.setDate(QDate(2026, 10, 5))
    selector = dialog.items.cellWidget(0, 0)
    selector.setCurrentIndex(selector.findData(product["id"]))
    price = dialog.items.cellWidget(0, 3)
    assert price.text() == "12345"
    price.setText("11111")
    dialog.date.setDate(QDate(2026, 10, 20))
    assert price.text() == "11111"  # A date edit is not consent to overwrite a quote.
    dialog.price_button.click()
    assert price.text() == str(base)
    dialog.date.setDate(QDate(2026, 10, 5))
    dialog.price_button.click()
    assert price.text() == "12345"
    dialog.save()
    assert service.get_document(dialog.saved_id)["items"][0]["unit_price"] == 12345
    assert next(p for p in service.list_products() if p["id"] == product["id"])["sale_price"] == base
    assert messages == []


def test_all_desktop_workflows_start_and_persist_with_socket_access_disabled(workspace, monkeypatch, tmp_path):
    service, messages = workspace

    def no_network(*_args, **_kwargs):
        raise AssertionError("Offline business workflow attempted network access")

    monkeypatch.setattr(socket, "socket", no_network)
    monkeypatch.setattr(socket, "create_connection", no_network)
    window = MainWindow(service, service.path)
    for key in window.nav_buttons:
        assert window.navigate(key)
    product = service.list_products()[0]
    today = date.today().isoformat()
    batch = service.save_stocktake([{"product_id": product["id"], "expected_qty": product["stock_qty"], "counted_qty": product["stock_qty"]}], "오프라인 실사 확인")
    member = service.save_member({"code": "M-OFFLINE", "name": "오프라인 회원"})
    service.adjust_member_points(member, 123, "수기 적립")
    service.save_promotion({"product_id": product["id"], "title": "오프라인 행사", "date_from": today, "date_to": today, "sale_price": 100})
    report = service.grouped_report(today, today, "monthly")
    target = tmp_path / "offline.xlsx"
    export_workbook(target, "오프라인", ["구분", "매출"], [[r["label"], r["sale_total"]] for r in report["rows"]])
    assert load_workbook(target).active.max_column == 2
    backup = tmp_path / "backup.db"
    service.backup(backup)
    reopened = POSService(backup)
    try:
        reopened.login("demo", "Demo-2026!")
        assert reopened.member_points(member) == 123
        assert reopened.get_stocktake(batch)["reason"] == "오프라인 실사 확인"
        assert reopened.get_effective_price(product["id"], today, "sale") == 100
        assert len(reopened.list_documents("sale")) == 7
    finally:
        reopened.close()
    assert messages == []
