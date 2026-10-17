"""Excel exports preserve numeric accounting data and literal user strings."""
from decimal import Decimal

from openpyxl import load_workbook

from seoyoung_pos.export import export_workbook


def test_export_text_cannot_become_excel_formulas(tmp_path):
    values = [
        '=HYPERLINK("https://example.invalid", "거래처")',
        "+SUM(1,2)", "-2+3", "@SUM(1,2)", "\t=1+1", "정상 한글 상품명",
    ]
    target = export_workbook(tmp_path / "거래처.xlsx", "공급사/목록", ["메모"], [[value] for value in values])
    workbook = load_workbook(target, data_only=False)
    sheet = workbook.active
    assert sheet.title == "공급사·목록"
    assert [sheet.cell(index + 2, 1).value for index in range(len(values))] == values
    assert all(sheet.cell(index + 2, 1).data_type == "s" for index in range(len(values)))
    workbook.close()


def test_export_numbers_remain_numeric_and_workbook_is_filterable(tmp_path):
    rows = [["국내산 삼겹살", 23500, Decimal("1.250"), -1250, 0]]
    target = export_workbook(tmp_path / "거래.xlsx", "매입전표", ["상품명", "금액", "수량", "조정액", "잔액"], rows)
    workbook = load_workbook(target)
    sheet = workbook.active
    assert [sheet.cell(2, column).value for column in range(2, 6)] == [23500, 1.25, -1250, 0]
    assert all(sheet.cell(2, column).data_type == "n" for column in range(2, 6))
    assert all("#,##0" in sheet.cell(2, column).number_format for column in range(2, 6))
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:E2"
    assert sheet["A1"].font.bold
    assert sheet.column_dimensions["A"].width >= 14
    workbook.close()


def test_empty_export_keeps_headers_and_excel_filter(tmp_path):
    target = export_workbook(tmp_path / "빈자료.xlsx", "빈 자료", ["상품코드", "상품명"], [])
    workbook = load_workbook(target)
    sheet = workbook.active
    assert sheet.max_row == 1
    assert [cell.value for cell in sheet[1]] == ["상품코드", "상품명"]
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:B1"
    workbook.close()
