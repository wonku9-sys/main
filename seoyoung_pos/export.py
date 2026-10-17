"""Excel export with typed numbers and literal (never formula) user text."""
from pathlib import Path
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def export_workbook(path, title, headers, rows):
    target = Path(path)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title[:31].replace("/", "·")
    sheet.append(list(headers))
    for row in rows:
        sheet.append(list(row))
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
            elif isinstance(cell.value, (int, float, Decimal)):
                cell.number_format = '#,##0.###;[Red]-#,##0.###'
                cell.alignment = Alignment(horizontal="right")
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="14263E")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 30
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, header in enumerate(headers, start=1):
        widest = max((len(str(sheet.cell(r, index).value or "")) for r in range(1, min(sheet.max_row, 250) + 1)), default=len(header))
        sheet.column_dimensions[get_column_letter(index)].width = min(50, max(14, widest * 1.6 + 3))
    workbook.save(target)
    return target
